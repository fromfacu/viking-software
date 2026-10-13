"""Capa de servicios — Viking Software BETA · Etapa 1.8 (versión 0.5.0).

Autor: Facundo David Barreto
Cliente: Peluquería Vikingos (Juliano — propietario)

Reglas de negocio del módulo. Novedades de la Etapa 1.5:
    * RegistroVentaService exige barbero (trazabilidad) y acepta cliente,
      método de pago y propina.
    * GastoService registra salidas de caja.
    * AdminService concentra las operaciones del panel Administración.
    * ExportadorExcel genera el libro institucional estilo JULIO.

Novedades de la Etapa 1.6 (0.3.0):
    * Métodos de pago administrables (tabla metodos_pago).
    * Comisión % por barbero con liquidación en el reporte (propinas 100 %
      del barbero). Todos los textos fijos del reporte en MAYÚSCULAS.

Novedades de la Etapa 1.7 (0.4.0):
    * Base de CLIENTES (nombre, apellido, celular único, mail opcional):
      la venta se puede vincular a un cliente registrado y el reporte
      suma la hoja CLIENTES con las estadísticas de cada uno,
      identificado por su celular.
    * AUTOCOMPLETADO de clientes en la caja (búsqueda parcial por
      nombre/apellido/celular/mail) y DETECTOR DE ERRORES DE TIPEO:
      si nombre + mail coinciden con una ficha existente pero el celular
      es otro, se consulta al operador antes de guardar.
    * MONEDA configurable (Administración → MÉTODOS Y MONEDA): el símbolo
      elegido (ej.: G$) se usa en TODA la lectura de importes, en pantalla
      y en el Excel exportado. No convierte valores: es solo lectura.

Novedades de la Etapa 1.8 (0.5.0):
    * ADELANTOS a barberos (Administración → BARBEROS): egreso de caja que
      se descuenta del NETO A COBRAR del barbero en su hoja individual.
    * Hoja BALANCE en el Excel: INGRESOS (ventas + propinas) − EGRESOS
      (gastos + adelantos) = BALANCE del período.
    * Clientes con FECHA DE NACIMIENTO opcional (base del CRM de la
      Etapa 2); viaja a la hoja CLIENTES del reporte.

Mensajes de 3 bloques (sintaxis algorítmica obligatoria):
    estado de la operación → causa traducida → acción táctica exacta.
    El operador jamás ve una traza de error técnica.
"""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .modelos import MONEDA_DEFECTO, MONEDAS, Adelanto, Gasto, TipoItem, Venta

log = logging.getLogger("viking")

# Tokens institucionales
AZUL_VK = "1D3557"
ROJO_VK = "E63946"
GRIS_FILA = "F2F2F2"
BLANCO = "FFFFFF"

FORMATO_PORCENTAJE = "0.00%"
DIAS_ES = ("LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES", "SÁBADO", "DOMINGO")


class ErrorValidacion(Exception):
    """El operador recibe mensaje de 3 bloques; el código jamás muestra trazas."""


class ClientePosibleDuplicado(Exception):
    """Se cargó un cliente con NOMBRE + MAIL iguales a uno existente pero
    con OTRO celular: probable error de tipeo del número. NO es un error
    de validación: la interfaz lo atrapa aparte y CONSULTA al operador
    (corregir la ficha existente / crear una nueva / cancelar).

    Atributos:
        existente:      la ficha ya registrada que coincide (Cliente).
        celular_nuevo:  el celular que se escribió y no coincide.
    """

    def __init__(self, existente, celular_nuevo: str):
        super().__init__(
            f"Nombre + mail ya registrados con otro celular "
            f"({existente.celular} <> {celular_nuevo})")
        self.existente = existente
        self.celular_nuevo = celular_nuevo


def formato_moneda(valor: float, simbolo: str = MONEDA_DEFECTO) -> str:
    """'G$ 45.500,00' — símbolo configurable + punto de miles y coma
    decimal (es-AR). El símbolo lo define la moneda del sistema."""
    texto = f"{valor:,.2f}"
    return f"{simbolo} " + texto.replace(",", "X").replace(".", ",").replace("X", ".")


def formato_celda_moneda(simbolo: str = MONEDA_DEFECTO) -> str:
    """Formato numérico Excel con el símbolo de la moneda del sistema:
    '"G$" #,##0.00'. Las comillas protegen símbolos multi-carácter."""
    return f'"{simbolo}" #,##0.00'


def _a_float(valor, defecto=0.0) -> float:
    try:
        return round(float(str(valor).replace(",", ".")), 2)
    except (TypeError, ValueError):
        return defecto


def _a_int(valor, defecto=0) -> int:
    try:
        return int(float(str(valor)))
    except (TypeError, ValueError):
        return defecto


# ======================================================================
# Registro de ventas (caja rápida)
# ======================================================================

class RegistroVentaService:
    """Valida, calcula y persiste una venta. Único autorizado a fechar.

    Matriz de decisión v4 (Etapa 1.7) — se evalúa en orden:
        C1 descripción no vacía · C2 barbero seleccionado ·
        C3 método de pago válido (existe en metodos_pago; se guarda con la
        escritura canónica) · C4 precio > 0 y cantidad > 0 · C5 propina >= 0
        C6 cliente válido (si se elige uno registrado, debe existir; su
        nombre se desnormaliza en la venta para el historial)
        Todo S → guardar; primera falla → mensaje de 3 bloques y nada se guarda.
    """

    def __init__(self, repo_ventas, repo_barberos, repo_metodos, repo_clientes):
        self.repo = repo_ventas
        self.barberos = repo_barberos
        self.metodos = repo_metodos
        self.clientes = repo_clientes

    def registrar(self, tipo: TipoItem, descripcion: str, precio, cantidad,
                  barbero_id, metodo_pago: str = "EFECTIVO",
                  cliente: str = "", propina=0.0, cliente_id=None) -> Venta:
        descripcion = (descripcion or "").strip()
        cliente = (cliente or "").strip()
        precio = _a_float(precio)
        cantidad = _a_int(cantidad)
        propina = _a_float(propina)
        barbero_id = _a_int(barbero_id)
        cliente_id = _a_int(cliente_id)

        if not descripcion:
            raise ErrorValidacion(
                "La venta no se guardó porque falta la descripción. "
                "Escribí o elegí un ítem y volvé a pulsar REGISTRAR VENTA.")
        if barbero_id <= 0 or self.barberos.obtener(barbero_id) is None:
            raise ErrorValidacion(
                "La venta no se guardó porque no hay un barbero seleccionado. "
                "Elegí quién atiende en el selector de arriba y volvé a pulsar REGISTRAR VENTA.")
        metodo_canonico = self.metodos.buscar_por_nombre(metodo_pago)
        if metodo_canonico is None:
            raise ErrorValidacion(
                "La venta no se guardó porque el método de pago no es válido. "
                "Elegí un método de la lista y volvé a pulsar REGISTRAR VENTA.")
        if precio <= 0 or cantidad <= 0:
            raise ErrorValidacion(
                "La venta no se guardó porque el precio y la cantidad deben ser mayores a cero. "
                "Corregí los importes y volvé a pulsar REGISTRAR VENTA.")
        if propina < 0:
            raise ErrorValidacion(
                "La venta no se guardó porque la propina no puede ser negativa. "
                "Corregí la propina (o dejala en 0) y volvé a pulsar REGISTRAR VENTA.")
        if cliente_id > 0:
            cli = self.clientes.obtener(cliente_id)
            if cli is None:
                raise ErrorValidacion(
                    "La venta no se guardó porque el cliente elegido ya no existe. "
                    "Volvé a elegirlo de la lista (o dejá MOSTRADOR) y pulsá REGISTRAR VENTA.")
            cliente = cli.nombre_completo

        venta = Venta(
            tipo=tipo, descripcion=descripcion, precio_unitario=precio,
            cantidad=cantidad, barbero_id=barbero_id, metodo_pago=metodo_canonico,
            cliente=cliente, cliente_id=cliente_id or None, propina=propina,
            fecha_hora=datetime.now().isoformat(timespec="seconds"))
        try:
            self.repo.guardar(venta)
        except Exception as exc:
            log.exception("Fallo al guardar la venta %s", venta.uuid)
            raise ErrorValidacion(
                "La venta no se guardó porque la base de datos no pudo escribir en disco. "
                "Cerrá y volvé a abrir Viking Software; si el problema sigue, avisá a soporte."
            ) from exc

        log.info("Venta registrada: %s x%s = %.2f (barbero_id=%s, cliente_id=%s, %s, %s)",
                 venta.descripcion, venta.cantidad, venta.total,
                 venta.barbero_id, venta.cliente_id, venta.metodo_pago, venta.uuid)
        return venta


# ======================================================================
# Gastos
# ======================================================================

class GastoService:
    """Registra salidas de dinero de la caja."""

    def __init__(self, repo_gastos):
        self.repo = repo_gastos

    def registrar(self, categoria: str, descripcion: str, monto) -> Gasto:
        categoria = (categoria or "").strip()
        descripcion = (descripcion or "").strip()
        monto = _a_float(monto)
        if not categoria:
            raise ErrorValidacion(
                "El gasto no se guardó porque falta la categoría. "
                "Elegí una categoría y volvé a pulsar REGISTRAR GASTO.")
        if not descripcion:
            raise ErrorValidacion(
                "El gasto no se guardó porque falta la descripción. "
                "Escribí en qué se gastó y volvé a pulsar REGISTRAR GASTO.")
        if monto <= 0:
            raise ErrorValidacion(
                "El gasto no se guardó porque el monto debe ser mayor a cero. "
                "Corregí el monto y volvé a pulsar REGISTRAR GASTO.")
        gasto = Gasto(categoria=categoria, descripcion=descripcion, monto=monto,
                      fecha_hora=datetime.now().isoformat(timespec="seconds"))
        self.repo.guardar(gasto)
        log.info("Gasto registrado: %s / %s = %.2f", categoria, descripcion, monto)
        return gasto


# ======================================================================
# Administración (barberos, catálogo, métodos de pago, clientes,
# movimientos, moneda, clave)
# ======================================================================

def _validar_comision(comision) -> float:
    """Regla única de comisión: número entre 0 y 100 (porcentaje)."""
    valor = _a_float(comision, defecto=-1)
    if valor < 0 or valor > 100:
        raise ErrorValidacion(
            "No se guardó porque la comisión debe ser un porcentaje entre 0 y 100 "
            "(por ejemplo: 10). Corregí el valor y volvé a intentar.")
    return valor


def _validar_fecha_nacimiento(fecha: str) -> str:
    """Regla única de nacimiento (Etapa 1.8): opcional; si se carga debe
    ser una fecha real en formato DD/MM/AAAA y no puede ser futura.
    Devuelve la fecha normalizada con ceros ("05/08/1990") o ""."""
    fecha = (fecha or "").strip()
    if not fecha:
        return ""
    try:
        dt = datetime.strptime(fecha, "%d/%m/%Y")
    except ValueError:
        raise ErrorValidacion(
            "No se guardó el cliente porque la fecha de nacimiento no es válida. "
            "Escribila como DD/MM/AAAA (por ejemplo 05/08/1990) o dejala "
            "vacía, es opcional.")
    if dt.date() > datetime.now().date():
        raise ErrorValidacion(
            "No se guardó el cliente porque la fecha de nacimiento es futura. "
            "Corregí la fecha (o dejala vacía, es opcional) y volvé a intentar.")
    return dt.strftime("%d/%m/%Y")


def _validar_datos_cliente(nombre: str, apellido: str, celular: str,
                           mail: str, fecha_nacimiento: str = "") -> tuple:
    """Regla única de cliente: nombre, apellido y celular obligatorios;
    mail opcional pero con pinta de mail si se carga; nacimiento opcional
    pero DD/MM/AAAA real y no futura si se carga. Devuelve los datos
    limpios o lanza ErrorValidacion con mensaje de 3 bloques."""
    nombre = (nombre or "").strip()
    apellido = (apellido or "").strip()
    celular = (celular or "").strip()
    mail = (mail or "").strip()
    if not nombre:
        raise ErrorValidacion(
            "No se guardó el cliente porque falta el nombre. "
            "Escribí el nombre y volvé a intentar.")
    if not apellido:
        raise ErrorValidacion(
            "No se guardó el cliente porque falta el apellido. "
            "Escribí el apellido y volvé a intentar.")
    if not celular:
        raise ErrorValidacion(
            "No se guardó el cliente porque falta el celular. "
            "El celular es el dato que identifica a cada cliente: "
            "escribilo y volvé a intentar.")
    if mail and ("@" not in mail or "." not in mail.split("@")[-1]):
        raise ErrorValidacion(
            "No se guardó el cliente porque el mail no parece válido. "
            "Corregí el mail (o dejalo vacío, es opcional) y volvé a intentar.")
    nacimiento = _validar_fecha_nacimiento(fecha_nacimiento)
    return nombre, apellido, celular, mail, nacimiento


class AdminService:
    """Operaciones del panel Administración. Traduce errores técnicos
    a mensajes de 3 bloques."""

    def __init__(self, repo_barberos, repo_catalogo, repo_ventas, repo_config,
                 repo_metodos, repo_clientes, repo_adelantos):
        self.barberos = repo_barberos
        self.catalogo = repo_catalogo
        self.ventas = repo_ventas
        self.config = repo_config
        self.metodos = repo_metodos
        self.clientes = repo_clientes
        self.adelantos = repo_adelantos

    # ---------------- Barberos ----------------
    def agregar_barbero(self, nombre: str, comision=50.0) -> None:
        if not (nombre or "").strip():
            raise ErrorValidacion("No se agregó el barbero porque falta el nombre. "
                                  "Escribí el nombre y volvé a intentar.")
        comision = _validar_comision(comision)
        try:
            self.barberos.agregar(nombre, comision)
        except Exception as exc:
            raise ErrorValidacion(
                f"No se agregó el barbero porque ya existe uno llamado '{nombre.strip()}'. "
                "Usá otro nombre o reactivá el existente.") from exc
        log.info("Barbero agregado: %s (comisión %.2f%%)", nombre.strip(), comision)

    def renombrar_barbero(self, barbero_id: int, nombre: str) -> None:
        if not (nombre or "").strip():
            raise ErrorValidacion("No se renombró el barbero porque el nombre está vacío. "
                                  "Escribí un nombre y volvé a intentar.")
        try:
            self.barberos.renombrar(barbero_id, nombre)
        except Exception as exc:
            raise ErrorValidacion(
                f"No se renombró porque ya existe otro barbero llamado '{nombre.strip()}'. "
                "Elegí otro nombre.") from exc
        log.info("Barbero %s renombrado a %s", barbero_id, nombre.strip())

    def set_comision_barbero(self, barbero_id: int, comision) -> None:
        """Define el % de ganancia del barbero sobre lo facturado (0–100).
        Las propinas no se tocan: siempre son 100 % del barbero."""
        if self.barberos.obtener(_a_int(barbero_id)) is None:
            raise ErrorValidacion(
                "No se guardó la comisión porque el barbero no existe. "
                "Tocá una fila de la tabla y volvé a intentar.")
        comision = _validar_comision(comision)
        self.barberos.set_comision(_a_int(barbero_id), comision)
        log.info("Barbero %s: comisión %.2f%%", barbero_id, comision)

    def eliminar_barbero(self, barbero_id: int) -> None:
        if self.barberos.tiene_ventas(barbero_id):
            raise ErrorValidacion(
                "No se puede eliminar este barbero porque tiene ventas registradas a su nombre. "
                "Usá 'Desactivar' para que deje de aparecer en la caja sin perder el historial.")
        self.barberos.eliminar(barbero_id)
        log.info("Barbero %s eliminado", barbero_id)

    # ---------------- Adelantos (Etapa 1.8) ----------------
    def registrar_adelanto(self, barbero_id, monto, nota: str = "") -> None:
        """Adelanto del dueño a un barbero: egreso de caja que se
        descuenta del NETO A COBRAR del barbero en su liquidación.
        La fecha la asigna el sistema (reloj), jamás el operador."""
        barbero_id = _a_int(barbero_id)
        barbero = self.barberos.obtener(barbero_id)
        if barbero_id <= 0 or barbero is None:
            raise ErrorValidacion(
                "No se registró el adelanto porque no hay un barbero seleccionado. "
                "Tocá una fila de la tabla de barberos y volvé a intentar.")
        monto = _a_float(monto)
        if monto <= 0:
            raise ErrorValidacion(
                "No se registró el adelanto porque el monto debe ser mayor a cero. "
                "Corregí el monto y volvé a intentar.")
        adelanto = Adelanto(barbero_id=barbero_id, monto=monto,
                            nota=(nota or "").strip(),
                            fecha_hora=datetime.now().isoformat(timespec="seconds"))
        try:
            self.adelantos.guardar(adelanto)
        except Exception as exc:
            log.exception("Fallo al guardar el adelanto %s", adelanto.uuid)
            raise ErrorValidacion(
                "El adelanto no se registró porque la base de datos no pudo escribir en disco. "
                "Cerrá y volvé a abrir Viking Software; si el problema sigue, avisá a soporte."
            ) from exc
        log.info("Adelanto registrado: %s recibe %.2f (%s)",
                 barbero.nombre, monto, adelanto.uuid)

    def eliminar_adelanto(self, adelanto_id) -> None:
        """Baja física de un adelanto cargado por error. El NETO A COBRAR
        del barbero se recalcula solo porque siempre sale de disco."""
        if self.adelantos.obtener(_a_int(adelanto_id)) is None:
            raise ErrorValidacion(
                "No se eliminó el adelanto porque ya no existe. "
                "Actualizá la lista tocando de nuevo al barbero y volvé a intentar.")
        self.adelantos.eliminar(_a_int(adelanto_id))
        log.info("Adelanto %s ELIMINADO desde Administración", adelanto_id)

    # ---------------- Catálogo ----------------
    def agregar_item(self, nombre: str, tipo: TipoItem, precio) -> None:
        if not (nombre or "").strip():
            raise ErrorValidacion("No se agregó el ítem porque falta el nombre. "
                                  "Escribí el nombre y volvé a intentar.")
        if _a_float(precio) <= 0:
            raise ErrorValidacion("No se agregó el ítem porque el precio debe ser mayor a cero. "
                                  "Corregí el precio y volvé a intentar.")
        try:
            self.catalogo.agregar(nombre, tipo, _a_float(precio))
        except Exception as exc:
            raise ErrorValidacion(
                f"No se agregó el ítem porque ya existe uno llamado '{nombre.strip()}'. "
                "Editá el existente o reactiválo.") from exc
        log.info("Ítem de catálogo agregado: %s (%s)", nombre.strip(), tipo.value)

    def actualizar_item(self, item_id: int, nombre: str, tipo: TipoItem, precio) -> None:
        if not (nombre or "").strip() or _a_float(precio) <= 0:
            raise ErrorValidacion(
                "No se guardó el ítem porque el nombre no puede estar vacío y el precio "
                "debe ser mayor a cero. Corregí los datos y volvé a intentar.")
        self.catalogo.actualizar(item_id, nombre, tipo, _a_float(precio))
        log.info("Ítem %s actualizado: %s", item_id, nombre.strip())

    def eliminar_item(self, item_id: int) -> None:
        """Baja física de un ítem del catálogo (las ventas históricas
        guardan la descripción como texto y no se modifican)."""
        if self.catalogo.obtener(_a_int(item_id)) is None:
            raise ErrorValidacion(
                "No se eliminó el ítem porque ya no existe en el catálogo. "
                "Actualizá la lista y volvé a intentar.")
        self.catalogo.eliminar(_a_int(item_id))
        log.info("Ítem %s eliminado del catálogo", item_id)

    # ---------------- Métodos de pago ----------------
    def agregar_metodo_pago(self, nombre: str) -> None:
        nombre_limpio = (nombre or "").strip()
        if not nombre_limpio:
            raise ErrorValidacion(
                "No se agregó el método de pago porque falta el nombre. "
                "Escribí el nombre (por ejemplo: BILLETERA VIRTUAL) y volvé a intentar.")
        try:
            self.metodos.agregar(nombre_limpio)
        except Exception as exc:
            raise ErrorValidacion(
                f"No se agregó el método de pago porque ya existe uno llamado "
                f"'{nombre_limpio}'. Revisá la lista y volvé a intentar.") from exc
        log.info("Método de pago agregado: %s", nombre_limpio)

    def eliminar_metodo_pago(self, metodo_id: int) -> None:
        """Baja física. Las ventas históricas conservan el nombre como
        texto, así que el historial no se pierde. Regla de protección:
        la caja siempre necesita al menos UN método de pago."""
        if self.metodos.obtener(_a_int(metodo_id)) is None:
            raise ErrorValidacion(
                "No se eliminó el método de pago porque ya no existe. "
                "Actualizá la lista y volvé a intentar.")
        if self.metodos.contar() <= 1:
            raise ErrorValidacion(
                "No se puede eliminar este método de pago porque es el único que queda. "
                "La caja necesita al menos un método para cobrar: agregá otro antes "
                "de eliminarlo.")
        self.metodos.eliminar(_a_int(metodo_id))
        log.info("Método de pago %s eliminado", metodo_id)

    # ---------------- Clientes ----------------
    def _detectar_posible_tipeo(self, nombre: str, celular: str, mail: str,
                                excluir_id: int | None = None) -> None:
        """Regla anti-tipeo del alta/edición de clientes: si NOMBRE + MAIL
        coinciden con una ficha existente pero el CELULAR es distinto,
        sospecha de un número mal tipeado y avisa con
        ClientePosibleDuplicado (la interfaz consulta al operador).
        Sin mail no hay dato para sospechar: la regla no aplica."""
        if not mail:
            return
        existente = self.clientes.buscar_por_nombre_y_mail(
            nombre, mail, excluir_id=excluir_id)
        if existente is not None and existente.celular != celular:
            raise ClientePosibleDuplicado(existente, celular)

    def agregar_cliente(self, nombre: str, apellido: str, celular: str,
                        mail: str = "", fecha_nacimiento: str = "",
                        forzar: bool = False) -> int:
        """Alta de cliente. Devuelve el id asignado. El celular es la
        clave única: no puede repetirse (sin distinción de mayúsculas).
        ``forzar=True`` crea la ficha aunque nombre + mail coincidan con
        otra (el operador ya confirmó que NO fue un error de tipeo)."""
        nombre, apellido, celular, mail, nacimiento = _validar_datos_cliente(
            nombre, apellido, celular, mail, fecha_nacimiento)
        if not forzar:
            self._detectar_posible_tipeo(nombre, celular, mail)
        try:
            cliente_id = self.clientes.agregar(nombre, apellido, celular, mail,
                                               nacimiento)
        except Exception as exc:
            raise ErrorValidacion(
                f"No se agregó el cliente porque ya existe uno con el celular "
                f"'{celular}'. Buscá ese cliente en la lista y editá sus datos "
                "si hace falta.") from exc
        log.info("Cliente agregado: %s %s (%s)", nombre, apellido, celular)
        return cliente_id

    def actualizar_cliente(self, cliente_id: int, nombre: str, apellido: str,
                           celular: str, mail: str = "",
                           fecha_nacimiento: str = "",
                           forzar: bool = False) -> None:
        if self.clientes.obtener(_a_int(cliente_id)) is None:
            raise ErrorValidacion(
                "No se guardó el cliente porque ya no existe. "
                "Tocá una fila de la tabla y volvé a intentar.")
        nombre, apellido, celular, mail, nacimiento = _validar_datos_cliente(
            nombre, apellido, celular, mail, fecha_nacimiento)
        if not forzar:
            self._detectar_posible_tipeo(nombre, celular, mail,
                                         excluir_id=_a_int(cliente_id))
        try:
            self.clientes.actualizar(_a_int(cliente_id), nombre, apellido,
                                     celular, mail, nacimiento)
        except Exception as exc:
            raise ErrorValidacion(
                f"No se guardó el cliente porque ya existe OTRO con el celular "
                f"'{celular}'. Verificá el número y volvé a intentar.") from exc
        log.info("Cliente %s actualizado: %s %s (%s)",
                 cliente_id, nombre, apellido, celular)

    def eliminar_cliente(self, cliente_id: int) -> None:
        """Baja física del cliente. Las ventas ya registradas conservan su
        nombre como texto (el historial y los reportes no cambian)."""
        if self.clientes.obtener(_a_int(cliente_id)) is None:
            raise ErrorValidacion(
                "No se eliminó el cliente porque ya no existe. "
                "Actualizá la lista y volvé a intentar.")
        self.clientes.eliminar(_a_int(cliente_id))
        log.info("Cliente %s eliminado (las ventas conservan su nombre)", cliente_id)

    # ---------------- Moneda ----------------
    def obtener_moneda(self) -> str:
        """Símbolo de la moneda del sistema (lectura de importes)."""
        simbolo = self.config.obtener("moneda_simbolo", MONEDA_DEFECTO)
        return simbolo if simbolo in MONEDAS else MONEDA_DEFECTO

    def set_moneda(self, simbolo: str) -> None:
        """Cambia el símbolo de moneda de TODO el sistema (pantalla y
        reportes). No convierte valores: es solo lectura de datos."""
        simbolo = (simbolo or "").strip()
        if simbolo not in MONEDAS:
            raise ErrorValidacion(
                "No se cambió la moneda porque el símbolo elegido no es válido. "
                f"Elegí una de la lista ({', '.join(MONEDAS)}) y volvé a intentar.")
        self.config.establecer("moneda_simbolo", simbolo)
        log.info("Moneda del sistema: %s", simbolo)

    # ---------------- Movimientos ----------------
    def editar_venta(self, venta_id: int, *, barbero_id: int, tipo: TipoItem,
                     descripcion: str, precio_unitario, cantidad,
                     cliente: str, metodo_pago, propina,
                     cliente_id=None) -> None:
        """Corrección de un movimiento (p. ej. barbero equivocado).
        ``metodo_pago`` es texto; se guarda tal cual lo elige el admin
        (puede ser un método histórico ya eliminado de la lista).
        Si se vincula un cliente registrado, su nombre se desnormaliza."""
        metodo_txt = str(metodo_pago or "").strip()
        cliente_id = _a_int(cliente_id)
        if not (descripcion or "").strip():
            raise ErrorValidacion("No se guardó la corrección porque falta la descripción. "
                                  "Completá la descripción y volvé a intentar.")
        if _a_int(barbero_id) <= 0 or self.barberos.obtener(_a_int(barbero_id)) is None:
            raise ErrorValidacion("No se guardó la corrección porque el barbero no es válido. "
                                  "Elegí un barbero de la lista y volvé a intentar.")
        if not metodo_txt:
            raise ErrorValidacion("No se guardó la corrección porque falta el método de pago. "
                                  "Elegí uno de la lista y volvé a intentar.")
        if _a_float(precio_unitario) <= 0 or _a_int(cantidad) <= 0:
            raise ErrorValidacion("No se guardó la corrección porque precio y cantidad deben "
                                  "ser mayores a cero. Corregí los importes y volvé a intentar.")
        if _a_float(propina) < 0:
            raise ErrorValidacion("No se guardó la corrección porque la propina no puede ser "
                                  "negativa. Corregila y volvé a intentar.")
        cliente_txt = (cliente or "").strip()
        if cliente_id > 0:
            cli = self.clientes.obtener(cliente_id)
            if cli is None:
                raise ErrorValidacion(
                    "No se guardó la corrección porque el cliente elegido ya no existe. "
                    "Volvé a elegirlo de la lista (o dejá MOSTRADOR) y volvé a intentar.")
            cliente_txt = cli.nombre_completo
        self.ventas.actualizar(
            venta_id, barbero_id=_a_int(barbero_id), tipo=tipo,
            descripcion=descripcion, precio_unitario=_a_float(precio_unitario),
            cantidad=_a_int(cantidad), cliente=cliente_txt,
            cliente_id=cliente_id or None,
            metodo_pago=metodo_txt, propina=_a_float(propina))
        log.info("Venta %s corregida desde Administración", venta_id)

    def eliminar_venta(self, venta_id: int) -> None:
        self.ventas.eliminar(venta_id)
        log.info("Venta %s ELIMINADA desde Administración", venta_id)

    # ---------------- Clave ----------------
    def cambiar_clave(self, clave_actual: str, clave_nueva: str) -> None:
        if not self.config.verificar_clave_admin(clave_actual):
            raise ErrorValidacion("No se cambió la clave porque la clave actual no coincide. "
                                  "Verificá la clave actual y volvé a intentar.")
        if len(clave_nueva or "") < 4:
            raise ErrorValidacion("No se cambió la clave porque la nueva debe tener al menos "
                                  "4 caracteres. Elegí una clave más larga.")
        self.config.cambiar_clave_admin(clave_nueva)
        log.info("Clave de administración actualizada")


# ======================================================================
# Exportación Excel — libro institucional estilo JULIO
# ======================================================================

class ExportadorExcel:
    """Genera el .xlsx del rango pedido (canal BAT-03) con cinco vistas:

    1. Hoja "Registro": tabla maestra completa (todos los campos).
    2. Hoja "Dashboard": resumen general, por barbero (con A LIQUIDAR),
       por método de pago, por tipo, gastos y neto del período.
    3. Hoja "Balance" (Etapa 1.8): INGRESOS (ventas + propinas) − EGRESOS
       (gastos + adelantos) = BALANCE del período, con el detalle de los
       adelantos entregados a cada barbero.
    4. Hoja "Clientes": estadísticas de cada cliente en el período,
       identificado por su CELULAR (operaciones, total gastado, propinas,
       ticket promedio, última visita y fecha de nacimiento).
    5. Una hoja por barbero ("Resumen <Nombre>"): bloques por día con
       Descripción/Monto/Propina y TOTAL por día con fórmulas, más el
       cierre con TOTAL FACTURADO y la LIQUIDACIÓN DEL BARBERO:
       comisión = % × facturado (sin propinas) + propinas 100 % suyas.
       El % queda en una celda editable: si cambia, el Excel recalcula.
       Etapa 1.8: la hoja lista los ADELANTOS del barbero y cierra con
       el NETO A COBRAR (liquidación − adelantos).

    Todos los importes usan el símbolo de la MONEDA DEL SISTEMA
    (config "moneda_simbolo"), leído al momento de exportar.
    """

    def __init__(self, repo_ventas, repo_gastos, repo_barberos, repo_config,
                 repo_adelantos, logo_path: Path, carpeta_salida: Path):
        self.repo = repo_ventas
        self.gastos = repo_gastos
        self.barberos = repo_barberos
        self.config = repo_config
        self.adelantos = repo_adelantos
        self.logo_path = Path(logo_path)
        self.carpeta_salida = Path(carpeta_salida)
        self.carpeta_salida.mkdir(parents=True, exist_ok=True)

    # --------------------------------------------------------------
    def _fmt(self) -> str:
        """Formato numérico de moneda con el símbolo vigente al exportar."""
        simbolo = self.config.obtener("moneda_simbolo", MONEDA_DEFECTO)
        return formato_celda_moneda(simbolo if simbolo in MONEDAS else MONEDA_DEFECTO)

    def exportar(self, fecha_desde: str, fecha_hasta: str) -> Path:
        filas = self.repo.ventas_entre_fechas(fecha_desde, fecha_hasta)
        if not filas:
            raise ErrorValidacion(
                "No se generó el reporte porque no hay ventas registradas en ese rango de fechas. "
                "Elegí otro rango y volvé a pulsar EXPORTAR DATOS.")

        nombre = f"reporte_ventas_{fecha_desde}_a_{fecha_hasta}.xlsx"
        destino = self.carpeta_salida / nombre
        try:
            comisiones = {b.nombre: b.comision
                          for b in self.barberos.listar(solo_activos=False)}
            adelantos = self.adelantos.listar_entre_fechas(fecha_desde,
                                                           fecha_hasta)
            wb = Workbook()
            self._hoja_registro(wb.active, filas, fecha_desde, fecha_hasta)
            self._hoja_dashboard(wb.create_sheet("Dashboard"), filas,
                                 fecha_desde, fecha_hasta, comisiones)
            self._hoja_balance(wb.create_sheet("Balance"), filas,
                               fecha_desde, fecha_hasta, adelantos)
            self._hoja_clientes(wb.create_sheet("Clientes"),
                                fecha_desde, fecha_hasta)
            # Un barbero tiene hoja propia si tuvo VENTAS o ADELANTOS en
            # el período (puede no haber vendido y aun así deber adelantos).
            barberos_hojas = sorted(set(self._barberos_con_ventas(filas))
                                    | {a[2] for a in adelantos})
            for barbero in barberos_hojas:
                hoja = wb.create_sheet(f"Resumen {barbero[:25]}")
                self._hoja_barbero(hoja, barbero, filas, comisiones.get(barbero),
                                   [a for a in adelantos if a[2] == barbero])
            wb.save(destino)
        except ErrorValidacion:
            raise
        except Exception as exc:
            log.exception("Fallo al exportar %s", nombre)
            raise ErrorValidacion(
                "El reporte no se generó porque Excel no pudo escribir el archivo. "
                "Cerrá cualquier reporte abierto en Excel y volvé a pulsar EXPORTAR DATOS."
            ) from exc
        log.info("Reporte exportado: %s (%s ventas)", destino, len(filas))
        return destino

    # --------------------------------------------------------------
    # utilidades internas
    # --------------------------------------------------------------
    @staticmethod
    def _barberos_con_ventas(filas) -> list[str]:
        vistos = []
        for f in filas:
            if f[2] not in vistos:
                vistos.append(f[2])
        return sorted(vistos)

    @staticmethod
    def _encabezado(hoja, titulos, fila, acento_ultima=True):
        fino = Side(style="thin", color="BFBFBF")
        for col, titulo in enumerate(titulos, start=1):
            c = hoja.cell(row=fila, column=col, value=titulo)
            c.fill = PatternFill("solid", fgColor=AZUL_VK)
            c.font = Font(bold=True, color=BLANCO)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = Border(left=fino, right=fino, top=fino, bottom=fino)
        if acento_ultima:
            hoja.cell(row=fila, column=len(titulos)).fill = PatternFill("solid", fgColor=ROJO_VK)

    def _logo_y_titulo(self, hoja, subtitulo: str) -> None:
        if self.logo_path.exists():
            img = XLImage(str(self.logo_path))
            # Altura fija 64 px conservando la proporción original del logo.
            if img.width and img.height:
                proporcion = img.width / img.height
                img.height = 64
                img.width = max(1, round(64 * proporcion))
            else:
                img.width, img.height = 64, 64
            hoja.add_image(img, "A1")
        hoja["C1"] = "VIKING SOFTWARE BETA"
        hoja["C1"].font = Font(bold=True, size=16, color=AZUL_VK)
        hoja["C2"] = subtitulo
        hoja["C2"].font = Font(size=12, color="444444")
        hoja["C3"] = f"GENERADO: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        hoja["C3"].font = Font(size=10, italic=True, color="666666")

    # --------------------------------------------------------------
    # 1) Tabla maestra
    # --------------------------------------------------------------
    def _hoja_registro(self, hoja, filas, desde, hasta) -> None:
        fmt = self._fmt()
        hoja.title = "Registro"
        self._logo_y_titulo(hoja, f"PELUQUERÍA VIKINGOS — REGISTRO DE VENTAS · {desde} A {hasta}")
        titulos = ("FECHA Y HORA", "BARBERO", "TIPO", "DESCRIPCIÓN", "CLIENTE",
                   "MÉTODO DE PAGO", "P. UNIT.", "CANT.", "SUBTOTAL",
                   "PROPINA", "TOTAL")
        self._encabezado(hoja, titulos, 5)

        fino = Side(style="thin", color="BFBFBF")
        borde = Border(left=fino, right=fino, top=fino, bottom=fino)
        r = 6
        for i, f in enumerate(filas):
            (vid, fecha_hora, barbero, tipo, descripcion, cliente, metodo,
             precio, cantidad, subtotal, propina, total) = f
            valores = (datetime.fromisoformat(fecha_hora).strftime("%d/%m/%Y %H:%M:%S"),
                       barbero, str(tipo).upper(), descripcion, cliente or "—",
                       str(metodo).upper(), precio, cantidad, subtotal, propina, total)
            for col, valor in enumerate(valores, start=1):
                c = hoja.cell(row=r, column=col, value=valor)
                c.border = borde
                if i % 2 == 1:
                    c.fill = PatternFill("solid", fgColor=GRIS_FILA)
                if col in (7, 9, 10, 11):
                    c.number_format = fmt
                    c.alignment = Alignment(horizontal="right")
                elif col == 8:
                    c.alignment = Alignment(horizontal="center")
            r += 1

        tot_sub = round(sum(f[9] for f in filas), 2)
        tot_prop = round(sum(f[10] for f in filas), 2)
        tot_gen = round(sum(f[11] for f in filas), 2)
        c = hoja.cell(row=r, column=8, value="GRAN TOTAL")
        c.font = Font(bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=AZUL_VK)
        c.alignment = Alignment(horizontal="right")
        for col, valor in ((9, tot_sub), (10, tot_prop), (11, tot_gen)):
            c = hoja.cell(row=r, column=col, value=valor)
            c.font = Font(bold=True, color=BLANCO)
            c.fill = PatternFill("solid", fgColor=ROJO_VK)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")

        # Anchos pensados para montos grandes con símbolo (ej.: G$ 120.000,00
        # = 13 caracteres): si la columna es angosta, Excel muestra ###.
        for col, ancho in zip(range(1, 12),
                              (19, 13, 11, 24, 14, 20, 15, 7, 15, 13, 15)):
            hoja.column_dimensions[get_column_letter(col)].width = ancho

    # --------------------------------------------------------------
    # 2) Dashboard
    # --------------------------------------------------------------
    def _hoja_dashboard(self, hoja, filas, desde, hasta, comisiones) -> None:
        fmt = self._fmt()
        self._logo_y_titulo(hoja, f"DASHBOARD GENERAL · {desde} A {hasta}")
        r = 5

        def titulo(txt):
            nonlocal r
            c = hoja.cell(row=r, column=1, value=txt)
            c.font = Font(bold=True, size=12, color=AZUL_VK)
            r += 1

        def par(rotulo, valor, moneda=True, negrita=False):
            nonlocal r
            hoja.cell(row=r, column=1, value=rotulo).font = Font(bold=negrita)
            c = hoja.cell(row=r, column=2, value=valor)
            if moneda:
                c.number_format = fmt
                c.alignment = Alignment(horizontal="right")
            if negrita:
                c.font = Font(bold=True)
            r += 1

        tot_sub = round(sum(f[9] for f in filas), 2)
        tot_prop = round(sum(f[10] for f in filas), 2)
        tot_gen = round(tot_sub + tot_prop, 2)
        gastos = self.gastos.listar_entre_fechas(desde, hasta)
        tot_gastos = round(sum(g[4] for g in gastos), 2)

        titulo("RESUMEN GENERAL DEL PERÍODO")
        par("VENTAS (SERVICIOS + PRODUCTOS)", tot_sub)
        par("PROPINAS", tot_prop)
        par("TOTAL COBRADO", tot_gen, negrita=True)
        par("GASTOS DEL PERÍODO", tot_gastos)
        par("NETO (COBRADO − GASTOS)", round(tot_gen - tot_gastos, 2), negrita=True)
        par("OPERACIONES REGISTRADAS", len(filas), moneda=False)
        r += 1

        titulo("POR BARBERO")
        self._encabezado(hoja, ("BARBERO", "OPERACIONES", "FACTURADO",
                                "PROPINAS", "TOTAL", "A LIQUIDAR"),
                         r, acento_ultima=False)
        r += 1
        for nombre in self._barberos_con_ventas(filas):
            fs = [f for f in filas if f[2] == nombre]
            fact = round(sum(f[9] for f in fs), 2)
            prop = round(sum(f[10] for f in fs), 2)
            hoja.cell(row=r, column=1, value=nombre).font = Font(bold=True)
            hoja.cell(row=r, column=2, value=len(fs)).alignment = Alignment(horizontal="center")
            for col, v in ((3, fact), (4, prop), (5, round(fact + prop, 2))):
                c = hoja.cell(row=r, column=col, value=v)
                c.number_format = fmt
                c.alignment = Alignment(horizontal="right")
            # Liquidación: comisión % sobre lo facturado (sin propinas)
            # + propinas, que son 100 % del barbero.
            comision = comisiones.get(nombre)
            c = hoja.cell(row=r, column=6)
            if comision is None:
                c.value = "—"
                c.alignment = Alignment(horizontal="center")
            else:
                c.value = round(fact * comision / 100 + prop, 2)
                c.number_format = fmt
                c.alignment = Alignment(horizontal="right")
            r += 1
        r += 1

        titulo("POR MÉTODO DE PAGO")
        self._encabezado(hoja, ("MÉTODO", "OPERACIONES", "TOTAL"), r, acento_ultima=False)
        r += 1
        metodos = sorted({str(f[6]).upper() for f in filas})
        for m in metodos:
            fs = [f for f in filas if str(f[6]).upper() == m]
            hoja.cell(row=r, column=1, value=m)
            hoja.cell(row=r, column=2, value=len(fs)).alignment = Alignment(horizontal="center")
            c = hoja.cell(row=r, column=3, value=round(sum(f[11] for f in fs), 2))
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
            r += 1
        r += 1

        titulo("POR TIPO")
        self._encabezado(hoja, ("TIPO", "OPERACIONES", "SUBTOTAL"), r, acento_ultima=False)
        r += 1
        for t in ("SERVICIO", "PRODUCTO"):
            fs = [f for f in filas if f[3] == t]
            if not fs:
                continue
            hoja.cell(row=r, column=1, value=t)
            hoja.cell(row=r, column=2, value=len(fs)).alignment = Alignment(horizontal="center")
            c = hoja.cell(row=r, column=3, value=round(sum(f[9] for f in fs), 2))
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
            r += 1
        r += 1

        titulo("GASTOS DEL PERÍODO")
        if gastos:
            self._encabezado(hoja, ("FECHA", "CATEGORÍA", "DESCRIPCIÓN", "MONTO"),
                             r, acento_ultima=False)
            r += 1
            for g in gastos:
                hoja.cell(row=r, column=1,
                          value=datetime.fromisoformat(g[1]).strftime("%d/%m/%Y %H:%M"))
                hoja.cell(row=r, column=2, value=g[2])
                hoja.cell(row=r, column=3, value=g[3])
                c = hoja.cell(row=r, column=4, value=g[4])
                c.number_format = fmt
                c.alignment = Alignment(horizontal="right")
                r += 1
            hoja.cell(row=r, column=3, value="TOTAL GASTOS").font = Font(bold=True)
            c = hoja.cell(row=r, column=4, value=tot_gastos)
            c.font = Font(bold=True)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
        else:
            hoja.cell(row=r, column=1, value="SIN GASTOS REGISTRADOS EN EL PERÍODO.")

        for col, ancho in zip(range(1, 7), (42, 16, 14, 14, 14, 14)):
            hoja.column_dimensions[get_column_letter(col)].width = ancho

    # --------------------------------------------------------------
    # 3) Balance — INGRESOS − EGRESOS del período (Etapa 1.8)
    # --------------------------------------------------------------
    def _hoja_balance(self, hoja, filas, desde, hasta, adelantos) -> None:
        """Caja del período en una mirada:
            INGRESOS  = ventas (subtotales) + propinas cobradas
            EGRESOS   = gastos + adelantos entregados a barberos
            BALANCE   = INGRESOS − EGRESOS
        Los adelantos son egreso de caja en el momento de entregarse;
        además se descuentan del NETO A COBRAR en la hoja de cada barbero
        (no se pierden: vuelven cuando se liquida)."""
        fmt = self._fmt()
        self._logo_y_titulo(hoja, f"BALANCE DEL PERÍODO · {desde} A {hasta}")
        r = 5

        def titulo(txt):
            nonlocal r
            c = hoja.cell(row=r, column=1, value=txt)
            c.font = Font(bold=True, size=12, color=AZUL_VK)
            r += 1

        def par(rotulo, valor, negrita=False):
            nonlocal r
            hoja.cell(row=r, column=1, value=rotulo).font = Font(bold=negrita)
            c = hoja.cell(row=r, column=2, value=valor)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
            if negrita:
                c.font = Font(bold=True)
            r += 1

        tot_sub = round(sum(f[9] for f in filas), 2)
        tot_prop = round(sum(f[10] for f in filas), 2)
        tot_ing = round(tot_sub + tot_prop, 2)
        tot_gastos = round(sum(g[4] for g in
                               self.gastos.listar_entre_fechas(desde, hasta)), 2)
        tot_adel = round(sum(a[3] for a in adelantos), 2)
        tot_egr = round(tot_gastos + tot_adel, 2)
        balance = round(tot_ing - tot_egr, 2)

        titulo("INGRESOS DEL PERÍODO")
        par("VENTAS (SERVICIOS + PRODUCTOS)", tot_sub)
        par("PROPINAS COBRADAS", tot_prop)
        par("TOTAL INGRESOS", tot_ing, negrita=True)
        r += 1

        titulo("EGRESOS DEL PERÍODO")
        par("GASTOS", tot_gastos)
        par("ADELANTOS A BARBEROS", tot_adel)
        par("TOTAL EGRESOS", tot_egr, negrita=True)
        r += 1

        # Cierre destacado: el número que le importa al dueño.
        c = hoja.cell(row=r, column=1, value="BALANCE (INGRESOS − EGRESOS)")
        c.font = Font(bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=AZUL_VK)
        c = hoja.cell(row=r, column=2, value=balance)
        c.font = Font(bold=True, color=BLANCO)
        c.fill = PatternFill("solid",
                             fgColor=AZUL_VK if balance >= 0 else ROJO_VK)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        r += 2

        # Detalle de adelantos: cada egreso queda auditado con fecha,
        # barbero y nota.
        titulo("DETALLE DE ADELANTOS DEL PERÍODO")
        if adelantos:
            self._encabezado(hoja, ("FECHA", "BARBERO", "MONTO", "NOTA"),
                             r, acento_ultima=False)
            r += 1
            fino = Side(style="thin", color="BFBFBF")
            borde = Border(left=fino, right=fino, top=fino, bottom=fino)
            for i, (aid, fecha, barbero, monto, nota) in enumerate(adelantos):
                valores = (datetime.fromisoformat(fecha).strftime("%d/%m/%Y %H:%M"),
                           barbero, monto, nota or "—")
                for col, valor in enumerate(valores, start=1):
                    c = hoja.cell(row=r, column=col, value=valor)
                    c.border = borde
                    if i % 2 == 1:
                        c.fill = PatternFill("solid", fgColor=GRIS_FILA)
                    if col == 3:
                        c.number_format = fmt
                        c.alignment = Alignment(horizontal="right")
                r += 1
            hoja.cell(row=r, column=2, value="TOTAL ADELANTOS").font = Font(bold=True)
            c = hoja.cell(row=r, column=3, value=tot_adel)
            c.font = Font(bold=True)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
            r += 2

        hoja.cell(row=r, column=1,
                  value="NOTA: LOS ADELANTOS SALEN DE CAJA HOY Y SE DESCUENTAN "
                        "DEL NETO A COBRAR DE CADA BARBERO EN SU HOJA DE "
                        "LIQUIDACIÓN. LAS PROPINAS SON 100 % DEL BARBERO."
                  ).font = Font(size=10, italic=True, color="666666")

        for col, ancho in zip(range(1, 5), (42, 18, 16, 40)):
            hoja.column_dimensions[get_column_letter(col)].width = ancho

    # --------------------------------------------------------------
    # 4) Clientes — estadísticas del período identificadas por celular
    # --------------------------------------------------------------
    def _hoja_clientes(self, hoja, desde, hasta) -> None:
        fmt = self._fmt()
        self._logo_y_titulo(hoja, f"CLIENTES · ESTADÍSTICAS DEL PERÍODO · {desde} A {hasta}")
        titulos = ("CLIENTE", "CELULAR", "MAIL", "NACIMIENTO", "OPERACIONES",
                   "TOTAL GASTADO", "PROPINAS", "TICKET PROMEDIO",
                   "ÚLTIMA VISITA")
        self._encabezado(hoja, titulos, 5, acento_ultima=False)

        stats = self.repo.estadisticas_clientes(desde, hasta)
        fino = Side(style="thin", color="BFBFBF")
        borde = Border(left=fino, right=fino, top=fino, bottom=fino)
        r = 6
        mostrador = None
        for i, (grupo, nombre, apellido, celular, mail, texto, ops,
                gastado, propinas, ultima, nacimiento) in enumerate(stats):
            if isinstance(grupo, str):
                if texto:
                    nombre_mostrar, cel_mostrar = texto, "—"
                    mail_mostrar, nac_mostrar = "—", "—"
                else:
                    mostrador = (ops, gastado, propinas, ultima)
                    continue          # la fila MOSTRADOR va siempre al final
            else:
                nombre_mostrar = f"{apellido}, {nombre}"
                cel_mostrar, mail_mostrar = celular, mail or "—"
                nac_mostrar = nacimiento or "—"
            self._fila_cliente(hoja, r, i, nombre_mostrar, cel_mostrar,
                               mail_mostrar, nac_mostrar, ops, gastado,
                               propinas, ultima, fmt, borde)
            r += 1
        if mostrador is not None:
            ops, gastado, propinas, ultima = mostrador
            self._fila_cliente(hoja, r, r, "SIN IDENTIFICAR (MOSTRADOR)",
                               "—", "—", "—", ops, gastado, propinas, ultima,
                               fmt, borde)
            r += 1

        # Cierre: totales del período (incluye mostrador)
        if stats:
            tot_ops = sum(f[6] for f in stats)
            tot_gas = round(sum(f[7] for f in stats), 2)
            tot_prop = round(sum(f[8] for f in stats), 2)
            c = hoja.cell(row=r, column=1, value="TOTAL GENERAL")
            c.font = Font(bold=True, color=BLANCO)
            c.fill = PatternFill("solid", fgColor=AZUL_VK)
            for col in (2, 3, 4, 8, 9):
                hoja.cell(row=r, column=col).fill = PatternFill("solid", fgColor=AZUL_VK)
            c = hoja.cell(row=r, column=5, value=tot_ops)
            c.font = Font(bold=True, color=BLANCO)
            c.fill = PatternFill("solid", fgColor=AZUL_VK)
            c.alignment = Alignment(horizontal="center")
            for col, valor in ((6, tot_gas), (7, tot_prop)):
                c = hoja.cell(row=r, column=col, value=valor)
                c.font = Font(bold=True, color=BLANCO)
                c.fill = PatternFill("solid", fgColor=ROJO_VK)
                c.number_format = fmt
                c.alignment = Alignment(horizontal="right")

        for col, ancho in zip(range(1, 10),
                              (28, 16, 26, 14, 13, 16, 13, 16, 18)):
            hoja.column_dimensions[get_column_letter(col)].width = ancho

    @staticmethod
    def _fila_cliente(hoja, r, i, nombre, celular, mail, nacimiento, ops,
                      gastado, propinas, ultima, fmt, borde) -> None:
        """Una fila de la hoja CLIENTES (cliente registrado o mostrador)."""
        valores = (nombre, celular, mail, nacimiento, ops, gastado, propinas,
                   round(gastado / ops, 2) if ops else 0,
                   datetime.fromisoformat(ultima).strftime("%d/%m/%Y %H:%M"))
        for col, valor in enumerate(valores, start=1):
            c = hoja.cell(row=r, column=col, value=valor)
            c.border = borde
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=GRIS_FILA)
            if col in (4, 5):
                c.alignment = Alignment(horizontal="center")
            elif col in (6, 7, 8):
                c.number_format = fmt
                c.alignment = Alignment(horizontal="right")

    # --------------------------------------------------------------
    # 5) Hoja por barbero (estilo libro JULIO + liquidación + adelantos)
    # --------------------------------------------------------------
    def _hoja_barbero(self, hoja, barbero: str, filas, comision,
                      adelantos=()) -> None:
        fmt = self._fmt()
        fs = [f for f in filas if f[2] == barbero]
        self._logo_y_titulo(hoja, f"RESUMEN INDIVIDUAL — {barbero}")

        # Agrupar por día
        dias: dict[str, list] = {}
        for f in fs:
            dias.setdefault(f[1][:10], []).append(f)

        r = 5
        totales_dia = []   # referencias a celdas de TOTAL de cada día
        for dia in sorted(dias):
            fecha = datetime.strptime(dia, "%Y-%m-%d").date()
            nombre_dia = f"{DIAS_ES[fecha.weekday()]} {fecha.strftime('%d-%m')}"
            c = hoja.cell(row=r, column=1, value=nombre_dia)
            c.font = Font(bold=True, color=BLANCO)
            c.fill = PatternFill("solid", fgColor=AZUL_VK)
            for col, t in ((2, "MONTO"), (3, "PROPINA")):
                c = hoja.cell(row=r, column=col, value=t)
                c.font = Font(bold=True, color=BLANCO)
                c.fill = PatternFill("solid", fgColor=ROJO_VK)
                c.alignment = Alignment(horizontal="center")
            r += 1

            ini = r
            for f in dias[dia]:
                hoja.cell(row=r, column=1,
                          value=f"{f[4]} — {f[5]}" if f[5] else f[4])
                c = hoja.cell(row=r, column=2, value=f[9])
                c.number_format = fmt
                c.alignment = Alignment(horizontal="right")
                if f[10] > 0:
                    c = hoja.cell(row=r, column=3, value=f[10])
                    c.number_format = fmt
                    c.alignment = Alignment(horizontal="right")
                r += 1
            fin = r - 1

            hoja.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
            c = hoja.cell(row=r, column=2, value=f"=SUM(B{ini}:B{fin})")
            c.font = Font(bold=True)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
            c2 = hoja.cell(row=r, column=3, value=f"=SUM(C{ini}:C{fin})")
            c2.font = Font(bold=True)
            c2.number_format = fmt
            c2.alignment = Alignment(horizontal="right")
            totales_dia.append((f"B{r}", f"C{r}"))
            r += 2  # línea en blanco entre días (como el libro JULIO)

        # Cierre del barbero. Si el barbero no tuvo ventas en el período
        # (puede tener hoja por sus ADELANTOS), los totales son 0 directo:
        # un =SUM() sin rangos sería una fórmula inválida.
        suma_montos = "+".join(ref for ref, _ in totales_dia)
        suma_propinas = "+".join(ref for _, ref in totales_dia)
        val_facturado = f"=SUM({suma_montos})" if suma_montos else 0
        val_propinas = f"=SUM({suma_propinas})" if suma_propinas else 0
        c = hoja.cell(row=r, column=1, value=f"TOTAL FACTURADO {str(barbero).upper()}")
        c.font = Font(bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=AZUL_VK)
        c = hoja.cell(row=r, column=2, value=val_facturado)
        c.font = Font(bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=AZUL_VK)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        fila_facturado = r
        r += 1
        hoja.cell(row=r, column=1, value="PROPINAS DEL PERÍODO").font = Font(bold=True)
        c = hoja.cell(row=r, column=2, value=val_propinas)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        fila_propinas = r
        r += 1
        hoja.cell(row=r, column=1, value="TOTAL CON PROPINAS").font = Font(bold=True)
        c = hoja.cell(row=r, column=2, value=f"=B{fila_facturado}+B{fila_propinas}")
        c.font = Font(bold=True)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        r += 2

        # LIQUIDACIÓN DEL BARBERO — solo si tiene comisión configurada.
        # Regla del negocio: la comisión se calcula SOLO sobre lo facturado
        # (subtotales); las propinas son 100 % del barbero y se suman aparte.
        if comision is not None:
            c = hoja.cell(row=r, column=1, value="LIQUIDACIÓN DEL BARBERO")
            c.font = Font(bold=True, color=BLANCO)
            c.fill = PatternFill("solid", fgColor=AZUL_VK)
            for col in (2, 3):
                hoja.cell(row=r, column=col).fill = PatternFill("solid", fgColor=AZUL_VK)
            r += 1

            hoja.cell(row=r, column=1, value="COMISIÓN PACTADA (% SOBRE LO FACTURADO)")
            c = hoja.cell(row=r, column=2, value=round(float(comision) / 100, 4))
            c.number_format = FORMATO_PORCENTAJE
            c.alignment = Alignment(horizontal="right")
            fila_pct = r            # celda editable: el Excel recalcula solo
            r += 1

            hoja.cell(row=r, column=1, value="COMISIÓN A PAGAR").font = Font(bold=True)
            c = hoja.cell(row=r, column=2,
                          value=f"=ROUND(B{fila_facturado}*B{fila_pct},2)")
            c.font = Font(bold=True)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
            fila_comision = r
            r += 1

            hoja.cell(row=r, column=1, value="PROPINAS (100 % DEL BARBERO)")
            c = hoja.cell(row=r, column=2, value=f"=B{fila_propinas}")
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
            fila_prop_liq = r
            r += 1

            c = hoja.cell(row=r, column=1, value="TOTAL A LIQUIDAR (COMISIÓN + PROPINAS)")
            c.font = Font(bold=True, color=BLANCO)
            c.fill = PatternFill("solid", fgColor=ROJO_VK)
            c = hoja.cell(row=r, column=2,
                          value=f"=B{fila_comision}+B{fila_prop_liq}")
            c.font = Font(bold=True, color=BLANCO)
            c.fill = PatternFill("solid", fgColor=ROJO_VK)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
            fila_liquidar = r
            r += 1

        # ADELANTOS DEL PERÍODO (Etapa 1.8): plata que el barbero ya
        # recibió del dueño; se descuenta de lo que hay que pagarle.
        r += 1
        c = hoja.cell(row=r, column=1, value="ADELANTOS DEL PERÍODO")
        c.font = Font(bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=AZUL_VK)
        for col in (2, 3):
            hoja.cell(row=r, column=col).fill = PatternFill("solid", fgColor=AZUL_VK)
        r += 1
        if adelantos:
            ini_adel = r
            for _aid, fecha, _nombre, monto, nota in adelantos:
                rotulo = datetime.fromisoformat(fecha).strftime("%d/%m/%Y")
                if nota:
                    rotulo += f" — {nota}"
                hoja.cell(row=r, column=1, value=rotulo)
                c = hoja.cell(row=r, column=2, value=monto)
                c.number_format = fmt
                c.alignment = Alignment(horizontal="right")
                r += 1
            fin_adel = r - 1
            hoja.cell(row=r, column=1, value="TOTAL ADELANTOS").font = Font(bold=True)
            c = hoja.cell(row=r, column=2, value=f"=SUM(B{ini_adel}:B{fin_adel})")
            c.font = Font(bold=True)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
        else:
            hoja.cell(row=r, column=1, value="SIN ADELANTOS EN EL PERÍODO")
            c = hoja.cell(row=r, column=2, value=0)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
        fila_adelantos = r
        r += 2

        # NETO A COBRAR: lo que hay que pagarle DESPUÉS de descontar los
        # adelantos. Fórmula viva: si se edita el % de comisión, el Excel
        # recalcula todo solo. Solo existe si el barbero tiene comisión.
        if comision is not None:
            c = hoja.cell(row=r, column=1,
                          value="NETO A COBRAR (LIQUIDACIÓN − ADELANTOS)")
            c.font = Font(bold=True, color=BLANCO)
            c.fill = PatternFill("solid", fgColor=ROJO_VK)
            for col in (2, 3):
                hoja.cell(row=r, column=col).fill = PatternFill("solid", fgColor=ROJO_VK)
            c = hoja.cell(row=r, column=2,
                          value=f"=B{fila_liquidar}-B{fila_adelantos}")
            c.font = Font(bold=True, color=BLANCO)
            c.fill = PatternFill("solid", fgColor=ROJO_VK)
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")

        hoja.column_dimensions["A"].width = 40
        hoja.column_dimensions["B"].width = 16
        hoja.column_dimensions["C"].width = 14
