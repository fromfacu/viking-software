"""Test manual de la Etapa 1.8 (0.5.0) — Viking Software BETA.

Autor: Facundo David Barreto
Cliente: Peluquería Vikingos (Juliano — propietario)

Verifica las capas de dominio, datos y servicios SIN abrir la interfaz:
registro con barbero/cliente/método/propina, panel de administración
(barberos con comisión, catálogo, movimientos, clave), métodos de pago
administrables, base de clientes (celular como clave única) con
AUTOCOMPLETADO (búsqueda parcial) y DETECTOR DE ERRORES DE TIPEO
(nombre + mail coinciden, celular no), moneda configurable, tema
claro/oscuro persistido, gastos, exportación multi-hoja estilo JULIO
con liquidación de comisiones y hoja CLIENTES, y las migraciones
automáticas v1→v5 y v2→v5.

Novedades Etapa 1.8 (0.5.0): la semilla de barberos carga SOLO al dueño
(Juliano — los empleados se agregan desde Administración); ADELANTOS del
dueño a los barberos (egreso de caja en la hoja BALANCE y descuento del
NETO A COBRAR en la hoja de cada barbero); FECHA DE NACIMIENTO opcional
en clientes (columna NACIMIENTO en la hoja CLIENTES); hoja BALANCE
(ingresos − egresos) en el Excel; modo claro pastel legible.

Ejecución (con el entorno virtual activo):
    python test_manual.py

Resultado esperado: "TODAS LAS PRUEBAS PASARON".
"""

import sqlite3
import tempfile
from contextlib import closing
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

from viking.base_datos import (ESQUEMA_BASE_SQL, ESQUEMA_V2_SQL,
                               AdelantoRepository, BarberoRepository,
                               CatalogoRepository, ClienteRepository,
                               ConexionDB, ConfigRepository, GastoRepository,
                               MetodoPagoRepository, VentaRepository)
from viking.modelos import MONEDAS, TipoItem
from viking.servicios import (AdminService, ClientePosibleDuplicado,
                              ErrorValidacion, ExportadorExcel,
                              GastoService, RegistroVentaService,
                              formato_celda_moneda, formato_moneda)

RAIZ = Path(__file__).resolve().parent


def _armar(ruta_db: Path):
    db = ConexionDB(ruta_db)
    db.inicializar()
    cat = CatalogoRepository(db)
    ven = VentaRepository(db)
    bar = BarberoRepository(db)
    gas = GastoRepository(db)
    cfg = ConfigRepository(db)
    met = MetodoPagoRepository(db)
    cli = ClienteRepository(db)
    ade = AdelantoRepository(db)
    reg = RegistroVentaService(ven, bar, met, cli)
    gsv = GastoService(gas)
    adm = AdminService(bar, cat, ven, cfg, met, cli, ade)
    exp = ExportadorExcel(ven, gas, bar, cfg, ade,
                          RAIZ / "assets" / "logo_vk.png",
                          ruta_db.parent / "exportaciones")
    return db, cat, ven, bar, gas, cfg, met, cli, ade, reg, gsv, adm, exp


def _rechaza(fn, *args, **kwargs) -> None:
    try:
        fn(*args, **kwargs)
        raise AssertionError(f"Debió rechazar: {fn.__name__}{args}{kwargs}")
    except ErrorValidacion:
        pass


def _rechaza_tipeo(fn, *args, **kwargs) -> None:
    """La operación debió frenar por SOSPECHA DE ERROR DE TIPEO
    (nombre + mail ya registrados con otro celular)."""
    try:
        fn(*args, **kwargs)
        raise AssertionError(f"Debió sospechar tipeo: {fn.__name__}{args}{kwargs}")
    except ClientePosibleDuplicado:
        pass


def _fila_por_rotulo(hoja, rotulo: str) -> int:
    """Fila (1-based) donde la columna A contiene exactamente `rotulo`."""
    for fila in hoja.iter_rows():
        if fila[0].value == rotulo:
            return fila[0].row
    raise AssertionError(f"No se encontró el rótulo '{rotulo}' en {hoja.title}")


def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        (db, cat, ven, bar, gas, cfg, met, cli, ade, reg, gsv, adm, exp) = \
            _armar(tmp / "viking_test.db")
        hoy = date.today().isoformat()

        # 1. Semillas: 7 ítems, el DUEÑO como único barbero de fábrica
        #    (comisión 50 %) y 4 métodos de pago; base de clientes vacía.
        #    Etapa 1.8: la semilla carga SOLO a Juliano (dueño); los
        #    empleados (como Alexander) se agregan desde Administración.
        assert len(cat.listar_activos()) == 7
        barberos = bar.listar()
        assert [b.nombre for b in barberos] == ["Juliano"], \
            "La semilla de fábrica debe cargar solo al dueño: Juliano"
        assert all(b.comision == 50.0 for b in barberos), \
            "La comisión de fábrica debe ser 50 %"
        juli = barberos[0]
        # Alexander es EMPLEADO: se carga desde Administración, como en
        # el local real (queda con la comisión de fábrica, 50 %)
        adm.agregar_barbero("Alexander")
        alex = next(b for b in bar.listar() if b.nombre == "Alexander")
        assert alex.comision == 50.0
        assert met.listar_nombres() == ["EFECTIVO", "TRANSFERENCIA",
                                        "TARJETA", "QR / Mercado Pago"]
        assert cli.listar() == [], "La base de clientes arranca vacía"

        # 2. Registro con barbero, cliente (texto libre legado), método de
        #    pago y propina. El método acepta cualquier escritura y guarda
        #    la canónica.
        reg.registrar(TipoItem.SERVICIO, "Corte Caballero", 8000, 1,
                      barbero_id=juli.id, metodo_pago="efectivo",  # minúsculas
                      cliente="Marcos", propina=500)
        reg.registrar(TipoItem.SERVICIO, "Perfilado de Barba", 5000, 2,
                      barbero_id=alex.id, metodo_pago="QR / Mercado Pago")
        reg.registrar(TipoItem.PRODUCTO, "Pomada", 6500, 1,
                      barbero_id=juli.id)
        reg.registrar(TipoItem.SERVICIO, "Coloración", 15000, 1,
                      barbero_id=alex.id, metodo_pago="transferencia",
                      propina=1500)

        # 3. Total del día desde disco: ventas + propinas
        esperado = 8000 + 500 + 5000 * 2 + 6500 + 15000 + 1500
        assert ven.total_del_dia(hoy) == esperado, \
            f"Total {ven.total_del_dia(hoy)} != {esperado}"
        filas = ven.ventas_del_dia(hoy)
        assert len(filas) == 4
        assert filas[0][2] == "Juliano" and filas[0][6] == "EFECTIVO", \
            "El método debió guardarse con la escritura canónica"
        assert filas[0][10] == 500 and filas[0][11] == 8500  # propina y total
        assert filas[3][6] == "TRANSFERENCIA"

        # 4. Validaciones: sin barbero, sin descripción, importes, propina,
        #    método inexistente, cliente inexistente
        _rechaza(reg.registrar, TipoItem.SERVICIO, "Corte", 8000, 1, 0)
        _rechaza(reg.registrar, TipoItem.SERVICIO, "Corte", 8000, 1, 99999)
        _rechaza(reg.registrar, TipoItem.SERVICIO, "", 8000, 1, juli.id)
        _rechaza(reg.registrar, TipoItem.SERVICIO, "Corte", 0, 1, juli.id)
        _rechaza(reg.registrar, TipoItem.SERVICIO, "Corte", 100, -1, juli.id)
        _rechaza(reg.registrar, TipoItem.SERVICIO, "Corte", 100, 1, juli.id,
                 propina=-1)
        _rechaza(reg.registrar, TipoItem.SERVICIO, "Corte", 100, 1, juli.id,
                 metodo_pago="BITCOIN")     # no existe en la lista
        _rechaza(reg.registrar, TipoItem.SERVICIO, "Corte", 100, 1, juli.id,
                 metodo_pago="")
        _rechaza(reg.registrar, TipoItem.SERVICIO, "Corte", 100, 1, juli.id,
                 cliente_id=99999)          # cliente inexistente
        assert len(ven.ventas_del_dia(hoy)) == 4, "Nada inválido debió guardarse"

        # 5. Admin — barberos: agregar con comisión, renombrar, set_comision,
        #    rango 0–100, baja protegida
        adm.agregar_barbero("Ricardo", comision=25)
        ricky = next(b for b in bar.listar() if b.nombre == "Ricardo")
        assert ricky.comision == 25.0
        adm.renombrar_barbero(ricky.id, "Ricardo R.")
        _rechaza(adm.agregar_barbero, "Juliano")                 # duplicado
        _rechaza(adm.agregar_barbero, "Otro", comision=101)      # fuera de rango
        _rechaza(adm.agregar_barbero, "Otro", comision=-5)
        _rechaza(adm.agregar_barbero, "Otro", comision="abc")    # no numérico
        adm.set_comision_barbero(alex.id, 10)                    # el caso real
        assert bar.obtener(alex.id).comision == 10.0
        adm.set_comision_barbero(alex.id, "10")                  # acepta texto
        _rechaza(adm.set_comision_barbero, alex.id, 150)
        _rechaza(adm.set_comision_barbero, alex.id, -1)
        _rechaza(adm.set_comision_barbero, alex.id, "")
        _rechaza(adm.set_comision_barbero, 99999, 10)            # no existe
        _rechaza(adm.eliminar_barbero, juli.id)                  # tiene ventas
        adm.eliminar_barbero(ricky.id)                           # sin ventas
        assert all(b.nombre != "Ricardo R." for b in bar.listar())

        # 6. Admin — catálogo: agregar, actualizar, desactivar, eliminar
        adm.agregar_item(nombre="Corte Niño", tipo=TipoItem.SERVICIO,
                         precio="6000")          # keyword, como llama la UI
        item = cat.buscar_por_nombre("Corte Niño")
        assert item is not None and item.precio_unitario == 6000
        _rechaza(adm.agregar_item, "Corte Niño", TipoItem.SERVICIO, 6000)
        _rechaza(adm.agregar_item, "X", TipoItem.SERVICIO, 0)
        todos = cat.listar_todos()
        iid = next(t[0] for t in todos if t[1] == "Corte Niño")
        adm.actualizar_item(iid, nombre="Corte Niño",
                            tipo=TipoItem.SERVICIO, precio=6500)
        assert cat.buscar_por_nombre("Corte Niño").precio_unitario == 6500
        cat.set_activo(iid, False)
        assert cat.buscar_por_nombre("Corte Niño").activo is False
        adm.eliminar_item(iid)
        assert cat.obtener(iid) is None
        _rechaza(adm.eliminar_item, iid)                         # ya no existe

        # 7. Regresión Inventario (bug 0.2.2): la UI debe llamar al servicio
        #    con parámetros POR NOMBRE. Y regresiones de la Etapa 1.7:
        #    cliente con AUTOCOMPLETADO en la caja, detector de errores de
        #    tipeo, sección MONEDA en la pestaña de métodos y botón de tema.
        fuente_ui = (RAIZ / "viking" / "interfaz.py").read_text(encoding="utf-8")
        assert "agregar_item(nombre=" in fuente_ui, \
            "PestañaInventario._agregar debe llamar con parámetros por nombre"
        assert "actualizar_item(iid, nombre=" in fuente_ui, \
            "PestañaInventario._actualizar debe llamar con parámetros por nombre"
        assert 'state="readonly", command=self._al_elegir_item' in fuente_ui, \
            "La lista de descripción debe ser de solo selección (readonly)"
        assert "self.combo_cliente = ComboAutocompletable(" in fuente_ui, \
            "El cliente de la caja debe tener autocompletado"
        assert "class ComboAutocompletable(ctk.CTkFrame)" in fuente_ui and \
            '"<KeyRelease>"' in fuente_ui, \
            "El autocompletado debe filtrar las sugerencias al tipear"
        # 0.4.3: la lista de sugerencias es PROPIA (tkinter puro) y jamás
        # usa el desplegable interno del ComboBox, que según la versión de
        # CTk es un MENÚ nativo que captura el teclado y frena la escritura.
        assert "_open_dropdown_menu" not in fuente_ui, \
            "Jamás abrir el desplegable interno del ComboBox (roba el teclado)"
        assert "tk.Listbox" in fuente_ui and "wm_overrideredirect" in fuente_ui, \
            "Las sugerencias deben ser una lista propia (Toplevel + Listbox)"
        assert "placeholder_text=MOSTRADOR" in fuente_ui, \
            "«— MOSTRADOR —» debe ser un texto guía (placeholder), no sólido"
        assert "self.combo_cliente.set(MOSTRADOR)" not in fuente_ui, \
            "El placeholder no debe escribirse como texto sólido en el campo"
        assert "self.campo_cliente" not in fuente_ui, \
            "El cliente ya no es un campo de texto libre plano en la caja"
        assert "ClientePosibleDuplicado" in fuente_ui and \
            "ERROR DE TIPEO" in fuente_ui, \
            "El alta/edición de clientes debe consultar errores de tipeo"
        assert "GUARDAR MONEDA" in fuente_ui and "MÉTODOS Y MONEDA" in fuente_ui, \
            "La pestaña de métodos debe incluir la sección MONEDA"
        assert "MODO CLARO" in fuente_ui and "aplicar_tema" in fuente_ui, \
            "La cabecera debe tener el botón MODO CLARO / MODO OSCURO"
        assert "self.combo_cliente.limpiar()" in fuente_ui, \
            "Tras registrar, el campo cliente debe volver al texto guía"
        assert "_guia_manual" in fuente_ui and "_mostrar_guia" in fuente_ui, \
            "El texto guía debe tener fallback manual (a prueba de versiones)"
        # 0.5.0: ADELANTOS en la pestaña BARBEROS, NACIMIENTO en CLIENTES
        # y el modo claro pastel con textos legibles (fichas y combos).
        assert "registrar_adelanto" in fuente_ui and "ADELANTOS" in fuente_ui, \
            "La pestaña BARBEROS debe incluir la carga de adelantos"
        assert "NACIMIENTO (OPC.):" in fuente_ui, \
            "La ficha del cliente debe pedir la fecha de nacimiento"
        assert "SEG_NO_SELECCIONADO" in fuente_ui and "BOTON_COMBO" in fuente_ui, \
            "El tema debe definir colores propios para fichas y combos"
        assert 'text_color="#FFFFFF"' in fuente_ui, \
            "Las fichas seleccionadas usan letra blanca sobre el azul VK"

        # 7b. Regresión 0.4.2 ("el programa no abre"): el punto de entrada
        #     debe capturar cualquier error fatal, registrarlo en el log y
        #     mostrar un cartel con instrucciones (pythonw no tiene consola).
        fuente_main = (RAIZ / "main.py").read_text(encoding="utf-8")
        assert "log.exception" in fuente_main and "ERROR FATAL" in fuente_main, \
            "main.py debe capturar y registrar los errores fatales al abrir"
        assert "messagebox" in fuente_main, \
            "main.py debe mostrar un cartel si el programa no pudo abrirse"

        # 8. Admin — métodos de pago: agregar, duplicado (sin distinción de
        #    mayúsculas), eliminar, protección del último
        adm.agregar_metodo_pago("Billetera Virtual")
        assert "Billetera Virtual" in met.listar_nombres()
        _rechaza(adm.agregar_metodo_pago, "billetera virtual")   # duplicado
        _rechaza(adm.agregar_metodo_pago, "efectivo")            # duplicado
        _rechaza(adm.agregar_metodo_pago, "")                    # vacío
        # El caso pedido: eliminar QR / Mercado Pago
        qr_id = next(i for i, n in met.listar() if n == "QR / Mercado Pago")
        adm.eliminar_metodo_pago(qr_id)
        assert "QR / Mercado Pago" not in met.listar_nombres()
        # Una venta ya registrada con QR conserva su texto en el historial
        assert ven.ventas_del_dia(hoy)[1][6] == "QR / Mercado Pago"
        # ...pero ya no se pueden registrar ventas nuevas con ese método
        _rechaza(reg.registrar, TipoItem.SERVICIO, "Corte", 8000, 1, juli.id,
                 metodo_pago="QR / Mercado Pago")
        _rechaza(adm.eliminar_metodo_pago, qr_id)                # ya no existe
        # Protección: nunca se puede eliminar el último método
        for i, n in list(met.listar()):
            if n != "EFECTIVO":
                adm.eliminar_metodo_pago(i)
        assert met.listar_nombres() == ["EFECTIVO"]
        _rechaza(adm.eliminar_metodo_pago, met.listar()[0][0])
        adm.agregar_metodo_pago("TARJETA")                       # se repone

        # 9. Clientes: alta, validaciones, unicidad del celular, edición,
        #    venta vinculada, baja con historial preservado.
        #    Etapa 1.8: la FECHA DE NACIMIENTO es opcional, se normaliza
        #    a DD/MM/AAAA y se valida que sea real y no futura.
        juan_id = adm.agregar_cliente("Juan", "Pérez", "3704111222",
                                      "juan@mail.com", "5/8/1990")
        ana_id = adm.agregar_cliente("Ana", "López", "3704333444")  # mail opc.
        assert cli.obtener(juan_id).nombre_completo == "Juan Pérez"
        assert cli.obtener(juan_id).fecha_nacimiento == "05/08/1990", \
            "El nacimiento debe guardarse normalizado a DD/MM/AAAA"
        assert cli.obtener(ana_id).mail == ""
        assert cli.obtener(ana_id).fecha_nacimiento == "", \
            "El nacimiento es opcional: vacío por defecto"
        _rechaza(adm.agregar_cliente, "", "Pérez", "3704111223")   # sin nombre
        _rechaza(adm.agregar_cliente, "Juan", "", "3704111223")    # sin apellido
        _rechaza(adm.agregar_cliente, "Juan", "Pérez", "")         # sin celular
        _rechaza(adm.agregar_cliente, "Juan", "Pérez", "3704111222")  # duplicado
        _rechaza(adm.agregar_cliente, "Otro", "Otro", "  3704111222 ")  # dup espacios
        _rechaza(adm.agregar_cliente, "Mala", "Mail", "3704999999",
                 "correo-sin-arroba")                              # mail inválido
        _rechaza(adm.agregar_cliente, "Mala", "Mail", "3704999999",
                 "juan@sinpunto")                                  # mail inválido
        # Nacimiento inválido: formato libre, fecha inexistente o futura
        _rechaza(adm.agregar_cliente, "Mala", "Fecha", "3704999997", "",
                 "ayer")                                           # formato inválido
        _rechaza(adm.agregar_cliente, "Mala", "Fecha", "3704999997", "",
                 "1990-08-15")                                     # formato inválido
        _rechaza(adm.agregar_cliente, "Mala", "Fecha", "3704999997", "",
                 "31/02/1990")                                     # fecha inexistente
        maniana = (date.today() + timedelta(days=1)).strftime("%d/%m/%Y")
        _rechaza(adm.agregar_cliente, "Mala", "Fecha", "3704999997", "",
                 maniana)                                          # fecha futura
        _rechaza(adm.actualizar_cliente, juan_id, "Juan", "Pérez",
                 "3704999888", "", "32/13/2000")                   # nac. inválido
        # Edición: cambia el celular de Juan (el que identifica al cliente)
        adm.actualizar_cliente(juan_id, "Juan", "Pérez", "3704999888", "",
                               "05/08/1990")
        assert cli.obtener(juan_id).celular == "3704999888"
        assert cli.obtener(juan_id).fecha_nacimiento == "05/08/1990", \
            "La edición debe conservar el nacimiento informado"
        assert cli.buscar_por_celular("3704111222") is None
        _rechaza(adm.actualizar_cliente, juan_id, "Juan", "Pérez",
                 "3704333444", "")                                 # celular de Ana
        _rechaza(adm.actualizar_cliente, 99999, "X", "Y", "3704000000")
        # Venta vinculada: el nombre se desnormaliza y viaja el cliente_id
        reg.registrar(TipoItem.SERVICIO, "Corte Caballero", 8000, 1,
                      barbero_id=juli.id, metodo_pago="EFECTIVO",
                      cliente_id=juan_id, propina=500)
        reg.registrar(TipoItem.SERVICIO, "Perfilado de Barba", 5000, 1,
                      barbero_id=alex.id, metodo_pago="TARJETA",
                      cliente_id=juan_id)
        reg.registrar(TipoItem.PRODUCTO, "Pomada", 6500, 1,
                      barbero_id=juli.id, metodo_pago="EFECTIVO",
                      cliente_id=ana_id)
        ventas_hoy = ven.ventas_del_dia(hoy)
        assert ventas_hoy[-3][5] == "Juan Pérez" and ventas_hoy[-2][5] == "Juan Pérez"
        assert ventas_hoy[-1][5] == "Ana López"
        with closing(db.abrir()) as con:
            cid_db = con.execute("SELECT cliente_id FROM ventas ORDER BY id DESC "
                                 "LIMIT 3").fetchall()
        assert cid_db[2][0] == juan_id and cid_db[0][0] == ana_id, \
            "cliente_id debe quedar guardado en la fila"
        v1_id = ventas_hoy[0][0]     # la venta vieja de texto libre ("Marcos")
        v7_id = ventas_hoy[-1][0]    # la venta de Ana
        # editar_venta puede vincular una venta de texto libre al cliente
        adm.editar_venta(v1_id, barbero_id=juli.id, tipo=TipoItem.SERVICIO,
                         descripcion="Corte Caballero", precio_unitario=8000,
                         cantidad=1, cliente="", cliente_id=juan_id,
                         metodo_pago="EFECTIVO", propina=500)
        assert ven.obtener(v1_id)[5] == "Juan Pérez", \
            "La corrección debe desnormalizar el nombre del cliente vinculado"
        # Baja: las ventas conservan el nombre como texto, cliente_id → NULL
        adm.eliminar_cliente(ana_id)
        assert cli.obtener(ana_id) is None
        assert ven.obtener(v7_id)[5] == "Ana López", \
            "La venta de un cliente eliminado conserva su nombre"
        with closing(db.abrir()) as con:
            assert con.execute("SELECT cliente_id FROM ventas WHERE id = ?",
                               (v7_id,)).fetchone()[0] is None
        _rechaza(adm.eliminar_cliente, ana_id)                     # ya no existe

        # 9b. Autocompletado y detector de errores de tipeo (pedido 0.4.0).
        #     Tres Richard cargados: tipear "rich" los sugiere a todos.
        adm.agregar_cliente("Richard", "Gómez", "3704000001")
        adm.agregar_cliente("Richard", "Paredes", "3704000002")
        adm.agregar_cliente("Richard", "Sosa", "3704000003",
                            "richard.sosa@mail.com")
        assert len(cli.buscar("rich")) == 3, \
            "Tipear 'rich' debe sugerir los 3 Richard cargados"
        assert len(cli.buscar("RICH")) == 3, \
            "La búsqueda no distingue mayúsculas"
        assert [c.apellido for c in cli.buscar("pared")] == ["Paredes"]
        assert [c.celular for c in cli.buscar("3704000001")] == ["3704000001"]
        assert [c.apellido for c in cli.buscar("richard sosa")] == ["Sosa"], \
            "La búsqueda combina nombre + apellido"
        assert cli.buscar("pérez")[0].nombre == "Juan"
        assert cli.buscar("zzz") == [], "Sin coincidencias: sin sugerencias"
        assert cli.buscar("100%_") == [], "Los comodines LIKE se escapan"

        # Detector de tipeo: NOMBRE + MAIL coinciden, el CELULAR no.
        carlos_id = adm.agregar_cliente("Carlos", "Herrera", "3704000010",
                                        "carlos@mail.com")
        _rechaza_tipeo(adm.agregar_cliente, "Carlos", "Herrera", "3704000011",
                       "CARLOS@MAIL.COM")      # mismo nombre+mail, otro celular
        # ...pero si el celular TAMBIÉN coincide, es el duplicado clásico
        _rechaza(adm.agregar_cliente, "Carlos", "Herrera", "3704000010",
                 "carlos@mail.com")
        # El operador confirma que es OTRA persona: forzar=True crea la ficha
        carlos2_id = adm.agregar_cliente("Carlos", "Herrera", "3704000011",
                                         "carlos@mail.com", forzar=True)
        assert cli.buscar_por_celular("3704000011") is not None
        # Sin mail no hay dato para sospechar: no aplica la regla
        adm.agregar_cliente("Carlos", "Herrera", "3704000012")
        # Edición: apuntar los datos hacia otra ficha existente también avisa
        _rechaza_tipeo(adm.actualizar_cliente, carlos2_id, "Carlos", "Herrera",
                       "3704000099", "carlos@mail.com")
        adm.actualizar_cliente(carlos2_id, "Carlos", "Herrera", "3704000099",
                               "carlos@mail.com", forzar=True)
        assert cli.obtener(carlos2_id).celular == "3704000099"
        # La búsqueda exacta nombre+mail no distingue mayúsculas y puede
        # excluir la ficha que se está editando
        assert cli.buscar_por_nombre_y_mail("carlos", "CARLOS@MAIL.COM") \
            is not None
        hallado = cli.buscar_por_nombre_y_mail("carlos", "carlos@mail.com",
                                               excluir_id=carlos_id)
        assert hallado is not None and hallado.id == carlos2_id
        assert cli.buscar_por_nombre_y_mail("carlos", "") is None, \
            "Sin mail no hay búsqueda nombre+mail"

        # 10. Admin — movimientos: corregir barbero equivocado y eliminar
        venta_mal = ven.ventas_del_dia(hoy)[1]  # Perfilado ×2 de Alexander
        adm.editar_venta(venta_mal[0], barbero_id=juli.id, tipo=TipoItem.SERVICIO,
                         descripcion="Perfilado de Barba", precio_unitario=5000,
                         cantidad=2, cliente="", metodo_pago="QR / Mercado Pago",
                         propina=0)
        corregida = ven.obtener(venta_mal[0])
        assert corregida[2] == "Juliano", "La corrección de barbero no se aplicó"
        assert corregida[6] == "QR / Mercado Pago", \
            "El admin puede conservar un método histórico al corregir"
        _rechaza(adm.editar_venta, venta_mal[0], barbero_id=juli.id,
                 tipo=TipoItem.SERVICIO, descripcion="X", precio_unitario=100,
                 cantidad=1, cliente="", metodo_pago="", propina=0)
        ultima = ven.ventas_del_dia(hoy)[3]  # Coloración: se elimina
        adm.eliminar_venta(ultima[0])
        assert ven.obtener(ultima[0]) is None

        # 11. Gastos
        gsv.registrar("Insumos", "Lavandina", 3500)
        _rechaza(gsv.registrar, "", "Lavandina", 3500)
        _rechaza(gsv.registrar, "Insumos", "Lavandina", 0)
        assert gas.total_entre_fechas(hoy, hoy) == 3500
        assert len(gas.listar_entre_fechas(hoy, hoy)) == 1

        # 11b. Adelantos a barberos (Etapa 1.8): plata que el dueño le da
        #      al barbero. Es un EGRESO DE CAJA (hoja BALANCE) y se
        #      descuenta del NETO A COBRAR en la liquidación del barbero.
        adm.registrar_adelanto(juli.id, 20000, "A cuenta de sueldo")
        adm.registrar_adelanto(alex.id, 5000)
        adel_juli = ade.listar_por_barbero(juli.id)
        assert len(adel_juli) == 1 and adel_juli[0][3] == 20000
        assert adel_juli[0][2] == "Juliano" and \
            adel_juli[0][4] == "A cuenta de sueldo"
        assert len(ade.listar_por_barbero(alex.id)) == 1
        assert ade.total_entre_fechas(hoy, hoy) == 25000
        assert len(ade.listar_entre_fechas(hoy, hoy)) == 2
        _rechaza(adm.registrar_adelanto, juli.id, 0)            # monto cero
        _rechaza(adm.registrar_adelanto, juli.id, -100)         # monto negativo
        _rechaza(adm.registrar_adelanto, juli.id, "abc")        # no numérico
        _rechaza(adm.registrar_adelanto, 99999, 1000)           # no existe
        _rechaza(adm.registrar_adelanto, 0, 1000)               # sin barbero
        assert ade.total_entre_fechas(hoy, hoy) == 25000, \
            "Ningún adelanto inválido debió guardarse"
        # Eliminar un adelanto cargado por error lo saca de la cuenta
        adm.registrar_adelanto(alex.id, 500, "error de prueba")
        aid_error = next(a[0] for a in ade.listar_por_barbero(alex.id)
                         if a[4] == "error de prueba")
        adm.eliminar_adelanto(aid_error)
        assert ade.obtener(aid_error) is None
        assert ade.total_entre_fechas(hoy, hoy) == 25000
        _rechaza(adm.eliminar_adelanto, aid_error)              # ya no existe

        # 12. Clave de administración
        assert cfg.verificar_clave_admin("vikingos")      # clave de fábrica
        assert not cfg.verificar_clave_admin("otra")
        adm.cambiar_clave("vikingos", "vikingos2026")
        assert cfg.verificar_clave_admin("vikingos2026")
        _rechaza(adm.cambiar_clave, "vikingos", "xxxx")   # actual incorrecta
        _rechaza(adm.cambiar_clave, "vikingos2026", "ab") # nueva muy corta

        # 13. Moneda del sistema (Etapa 1.7): símbolo en texto y en Excel
        assert adm.obtener_moneda() == "$", "La moneda de fábrica es $"
        assert formato_moneda(45500, "$") == "$ 45.500,00"
        _rechaza(adm.set_moneda, "")
        _rechaza(adm.set_moneda, "BTC")                       # fuera del catálogo
        for s in MONEDAS:                                     # todas válidas
            adm.set_moneda(s)
            assert adm.obtener_moneda() == s
        adm.set_moneda("G$")                                  # el caso pedido
        assert adm.obtener_moneda() == "G$"
        assert formato_moneda(45500, adm.obtener_moneda()) == "G$ 45.500,00"
        assert formato_celda_moneda("G$") == '"G$" #,##0.00'

        # 14. Tema claro/oscuro: se guarda en config y persiste
        assert cfg.obtener("tema", "oscuro") == "oscuro"
        cfg.establecer("tema", "claro")
        assert cfg.obtener("tema", "oscuro") == "claro"
        cfg.establecer("tema", "oscuro")

        # 15. Exportación multi-hoja estilo JULIO con liquidación y CLIENTES.
        #     Se exporta con la moneda del sistema en G$ para verificar que
        #     el símbolo elegido viaja al Excel.
        reg.registrar(TipoItem.SERVICIO, "Coloración", 15000, 1,
                      barbero_id=alex.id, metodo_pago="EFECTIVO",
                      propina=1500)
        destino = exp.exportar(hoy, hoy)
        assert destino.exists() and destino.stat().st_size > 4000
        libro = load_workbook(destino)
        assert "Registro" in libro.sheetnames
        assert "Dashboard" in libro.sheetnames
        assert "Balance" in libro.sheetnames, "Debe existir la hoja Balance"
        assert "Clientes" in libro.sheetnames, "Debe existir la hoja Clientes"
        assert "Resumen Juliano" in libro.sheetnames
        assert "Resumen Alexander" in libro.sheetnames
        reg_hoja = libro["Registro"]
        encabezados = [reg_hoja.cell(row=5, column=c).value for c in range(1, 12)]
        assert encabezados == ["FECHA Y HORA", "BARBERO", "TIPO", "DESCRIPCIÓN",
                               "CLIENTE", "MÉTODO DE PAGO", "P. UNIT.", "CANT.",
                               "SUBTOTAL", "PROPINA", "TOTAL"]
        # La moneda elegida (G$) se refleja en el formato de las celdas
        assert '"G$"' in reg_hoja.cell(row=6, column=9).number_format, \
            "El Excel debe usar el símbolo de la moneda del sistema"
        assert '"G$"' in reg_hoja.cell(row=6, column=11).number_format

        # Hoja CLIENTES: estadísticas por cliente identificado por celular.
        # Etapa 1.8: nueva columna NACIMIENTO (DD/MM/AAAA) tras MAIL.
        hoja_cli = libro["Clientes"]
        encabezados_cli = [hoja_cli.cell(row=5, column=c).value
                           for c in range(1, 10)]
        assert encabezados_cli == ["CLIENTE", "CELULAR", "MAIL", "NACIMIENTO",
                                   "OPERACIONES", "TOTAL GASTADO", "PROPINAS",
                                   "TICKET PROMEDIO", "ÚLTIMA VISITA"]
        fila_juan = None
        fila_ana = None
        fila_mostrador = None
        for fila in hoja_cli.iter_rows(min_row=6):
            if fila[1].value == "3704999888":
                fila_juan = fila
            if fila[0].value == "Ana López":
                fila_ana = fila
            if fila[0].value == "SIN IDENTIFICAR (MOSTRADOR)":
                fila_mostrador = fila
        assert fila_juan is not None, "Juan debe figurar identificado por celular"
        # Juan: 3 operaciones (Corte 8000 reasignado + Corte 8000 + Perfilado
        # 5000), propinas 1000, ticket promedio 7000
        assert fila_juan[0].value == "Pérez, Juan"
        assert fila_juan[3].value == "05/08/1990", \
            "El nacimiento del cliente debe viajar a la hoja CLIENTES"
        assert fila_juan[4].value == 3
        assert fila_juan[5].value == 21000
        assert fila_juan[6].value == 1000
        assert fila_juan[7].value == 7000.0        # ticket promedio
        assert '"G$"' in fila_juan[5].number_format
        # La venta de la cliente eliminada conserva su nombre como grupo propio
        assert fila_ana is not None, "La baja de un cliente no borra su historial"
        assert fila_ana[1].value == "—" and fila_ana[4].value == 1
        assert fila_ana[3].value == "—", \
            "Un cliente sin nacimiento cargado muestra «—»"
        assert fila_mostrador is not None, "Debe existir la fila MOSTRADOR"
        # Mostrador puro: Perfilado ×2 reasignado + Pomada + Coloración final
        assert fila_mostrador[4].value == 3
        r_total_cli = _fila_por_rotulo(hoja_cli, "TOTAL GENERAL")
        assert hoja_cli.cell(row=r_total_cli, column=5).value == \
            len(ven.ventas_del_dia(hoy))

        # Hoja BALANCE (Etapa 1.8): la caja del período en una mirada.
        #   INGRESOS = ventas + propinas cobradas
        #   EGRESOS  = gastos + adelantos entregados a barberos
        #   BALANCE  = INGRESOS − EGRESOS
        hoja_bal = libro["Balance"]
        ventas_periodo = ven.ventas_del_dia(hoy)
        tot_sub = round(sum(f[9] for f in ventas_periodo), 2)
        tot_prop = round(sum(f[10] for f in ventas_periodo), 2)
        r_vtas = _fila_por_rotulo(hoja_bal, "VENTAS (SERVICIOS + PRODUCTOS)")
        r_prp = _fila_por_rotulo(hoja_bal, "PROPINAS COBRADAS")
        r_ing = _fila_por_rotulo(hoja_bal, "TOTAL INGRESOS")
        r_gas = _fila_por_rotulo(hoja_bal, "GASTOS")
        r_ade = _fila_por_rotulo(hoja_bal, "ADELANTOS A BARBEROS")
        r_egr = _fila_por_rotulo(hoja_bal, "TOTAL EGRESOS")
        r_bal = _fila_por_rotulo(hoja_bal, "BALANCE (INGRESOS − EGRESOS)")
        assert hoja_bal.cell(row=r_vtas, column=2).value == tot_sub
        assert hoja_bal.cell(row=r_prp, column=2).value == tot_prop
        assert hoja_bal.cell(row=r_ing, column=2).value == tot_sub + tot_prop
        assert hoja_bal.cell(row=r_gas, column=2).value == 3500
        assert hoja_bal.cell(row=r_ade, column=2).value == 25000, \
            "Los adelantos son un egreso de caja en la hoja BALANCE"
        assert hoja_bal.cell(row=r_egr, column=2).value == 3500 + 25000
        assert hoja_bal.cell(row=r_bal, column=2).value == \
            tot_sub + tot_prop - 3500 - 25000, \
            "BALANCE = INGRESOS − EGRESOS"
        assert '"G$"' in hoja_bal.cell(row=r_bal, column=2).number_format
        # Cada adelanto queda auditado: fecha, barbero, monto y nota
        _fila_por_rotulo(hoja_bal, "DETALLE DE ADELANTOS DEL PERÍODO")
        detalle = [(f[1].value, f[2].value) for f in hoja_bal.iter_rows()
                   if f[1].value in ("Juliano", "Alexander")]
        assert sorted(detalle) == [("Alexander", 5000), ("Juliano", 20000)]

        # Hoja del barbero: fórmulas =SUM como el libro JULIO + liquidación
        resumen = libro["Resumen Alexander"]
        formulas = [c.value for fila in resumen.iter_rows() for c in fila
                    if isinstance(c.value, str) and c.value.startswith("=SUM")]
        assert formulas, "La hoja del barbero no tiene fórmulas =SUM"

        # Liquidación de Alexander (comisión 10 %, propinas 100 % suyas):
        r_fact = _fila_por_rotulo(resumen, "TOTAL FACTURADO ALEXANDER")
        r_prop = _fila_por_rotulo(resumen, "PROPINAS DEL PERÍODO")
        r_pct = _fila_por_rotulo(resumen, "COMISIÓN PACTADA (% SOBRE LO FACTURADO)")
        r_com = _fila_por_rotulo(resumen, "COMISIÓN A PAGAR")
        r_p100 = _fila_por_rotulo(resumen, "PROPINAS (100 % DEL BARBERO)")
        r_liq = _fila_por_rotulo(resumen, "TOTAL A LIQUIDAR (COMISIÓN + PROPINAS)")
        _fila_por_rotulo(resumen, "LIQUIDACIÓN DEL BARBERO")
        # El % queda en una celda editable en formato porcentaje (0.10 = 10 %)
        celda_pct = resumen.cell(row=r_pct, column=2)
        assert celda_pct.value == 0.10 and celda_pct.number_format == "0.00%"
        # COMISIÓN = facturado × % (SIN propinas), como fórmula viva
        assert resumen.cell(row=r_com, column=2).value == \
            f"=ROUND(B{r_fact}*B{r_pct},2)"
        # Propinas: referencia directa a las propinas del período
        assert resumen.cell(row=r_p100, column=2).value == f"=B{r_prop}"
        # Total a liquidar = comisión + propinas
        assert resumen.cell(row=r_liq, column=2).value == \
            f"=B{r_com}+B{r_p100}"
        # TOTAL CON PROPINAS sigue apuntando a facturado + propinas
        r_tcp = _fila_por_rotulo(resumen, "TOTAL CON PROPINAS")
        assert resumen.cell(row=r_tcp, column=2).value == \
            f"=B{r_fact}+B{r_prop}"

        # Dashboard: columna A LIQUIDAR con la cuenta ya hecha
        dash = libro["Dashboard"]
        r_pb = _fila_por_rotulo(dash, "POR BARBERO")
        encabezado_pb = [dash.cell(row=r_pb + 1, column=c).value
                         for c in range(1, 7)]
        assert encabezado_pb == ["BARBERO", "OPERACIONES", "FACTURADO",
                                 "PROPINAS", "TOTAL", "A LIQUIDAR"]
        encontrados = {}
        for fila in dash.iter_rows(min_row=r_pb + 2):
            nombre = fila[0].value
            if nombre in ("Alexander", "Juliano"):
                encontrados[nombre] = (fila[2].value, fila[3].value,
                                       fila[5].value)
        # Alexander: facturado 5000 (Perfilado de Juan) + 15000 (Coloración)
        # = 20000, propinas 1500, liquidar 2000 + 1500 = 3500
        assert encontrados["Alexander"] == (20000, 1500, 3500.0)
        # Juliano: facturado 8000 (Corte de Juan) + 10000 (Perfilado ×2
        # reasignado) + 6500 (Pomada) + 8000 (Corte de Juan) + 6500 (Pomada
        # de Ana) = 39000, propinas 1000, liquidar 19500 + 1000 = 20500
        assert encontrados["Juliano"] == (39000, 1000, 20500.0)

        # Hoja del barbero: ADELANTOS DEL PERÍODO y NETO A COBRAR
        # (liquidación − adelantos), como fórmulas vivas del Excel
        resumen_j = libro["Resumen Juliano"]
        r_liq_j = _fila_por_rotulo(resumen_j,
                                   "TOTAL A LIQUIDAR (COMISIÓN + PROPINAS)")
        r_ade_j = _fila_por_rotulo(resumen_j, "TOTAL ADELANTOS")
        r_net_j = _fila_por_rotulo(resumen_j,
                                   "NETO A COBRAR (LIQUIDACIÓN − ADELANTOS)")
        assert r_liq_j < r_ade_j < r_net_j, \
            "El NETO A COBRAR va después de la liquidación y los adelantos"
        assert resumen_j.cell(row=r_net_j, column=2).value == \
            f"=B{r_liq_j}-B{r_ade_j}", "NETO = LIQUIDACIÓN − ADELANTOS"
        # El adelanto de Juliano (20000) figura con su nota en su hoja
        rotulos_adel = [resumen_j.cell(row=r, column=1).value
                        for r in range(r_liq_j + 1, r_ade_j)]
        assert any("A CUENTA DE SUELDO" in str(rotulo).upper()
                   for rotulo in rotulos_adel)
        # Alexander también tiene adelanto (5000): misma estructura
        r_ade_a = _fila_por_rotulo(resumen, "TOTAL ADELANTOS")
        assert resumen.cell(row=r_ade_a, column=2).value is not None
        _fila_por_rotulo(resumen, "NETO A COBRAR (LIQUIDACIÓN − ADELANTOS)")

        _rechaza(exp.exportar, "1999-01-01", "1999-01-02")  # rango vacío

        # 16. Migración v1 → v5: una base vieja de la Etapa 1 no se rompe
        vieja = tmp / "viking_v1.db"
        with closing(sqlite3.connect(vieja)) as con:
            with con:
                con.execute("""CREATE TABLE items_catalogo (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    nombre TEXT NOT NULL UNIQUE,
                    tipo TEXT NOT NULL, precio_unitario REAL NOT NULL,
                    activo INTEGER NOT NULL DEFAULT 1)""")
                con.execute("""CREATE TABLE ventas (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, uuid TEXT NOT NULL UNIQUE,
                    fecha_hora TEXT NOT NULL, tipo TEXT NOT NULL,
                    descripcion TEXT NOT NULL, precio_unitario REAL NOT NULL,
                    cantidad INTEGER NOT NULL, subtotal REAL NOT NULL,
                    sincronizado INTEGER NOT NULL DEFAULT 0,
                    creado_en TEXT NOT NULL DEFAULT (datetime('now')),
                    actualizado_en TEXT NOT NULL DEFAULT (datetime('now')))""")
                con.execute("INSERT INTO items_catalogo (nombre, tipo, precio_unitario) "
                            "VALUES ('Corte Caballero', 'SERVICIO', 8000)")
                con.execute("INSERT INTO ventas (uuid, fecha_hora, tipo, descripcion, "
                            "precio_unitario, cantidad, subtotal) VALUES "
                            "('vieja-1', ?, 'SERVICIO', 'Corte Caballero', 8000, 1, 8000)",
                            (f"{hoy}T10:00:00",))
                con.execute("PRAGMA user_version = 1")

        db2 = ConexionDB(vieja)
        db2.inicializar()  # migra sola al arrancar
        ven2 = VentaRepository(db2)
        bar2 = BarberoRepository(db2)  # siembra solo a Juliano (el dueño)
        met2 = MetodoPagoRepository(db2)
        cli2 = ClienteRepository(db2)
        filas2 = ven2.ventas_del_dia(hoy)
        assert len(filas2) == 1, "La venta vieja debió conservarse"
        assert filas2[0][2] == "—", "La venta vieja queda sin barbero asignado"
        assert ven2.total_del_dia(hoy) == 8000
        assert all(b.comision == 50.0 for b in bar2.listar())
        assert len(met2.listar()) == 4, "Los métodos de pago se siembran al migrar"
        with closing(db2.abrir()) as con:
            assert con.execute("PRAGMA user_version").fetchone()[0] == 5
            tablas = {f[0] for f in con.execute(
                "SELECT name FROM sqlite_master WHERE type='table'")}
            assert "clientes" in tablas, "La migración debe crear la tabla clientes"
            assert "adelantos" in tablas, "La migración debe crear adelantos"
            cols_c = {f[1] for f in con.execute("PRAGMA table_info(clientes)")}
            assert "fecha_nacimiento" in cols_c, \
                "La migración debe agregar clientes.fecha_nacimiento"
            cols_v = {f[1] for f in con.execute("PRAGMA table_info(ventas)")}
            assert "cliente_id" in cols_v, "La migración debe crear ventas.cliente_id"
        reg2 = RegistroVentaService(ven2, bar2, met2, cli2)
        juli2 = next(b for b in bar2.listar() if b.nombre == "Juliano")
        reg2.registrar(TipoItem.SERVICIO, "Corte Caballero", 8000, 1,
                       barbero_id=juli2.id, propina=1000)
        assert ven2.total_del_dia(hoy) == 17000

        # 17. Migración v2 → v5: base de la Etapa 1.5 (barberos SIN comisión,
        #     sin tabla de métodos ni clientes) recibe todo lo nuevo
        v2 = tmp / "viking_v2.db"
        with closing(sqlite3.connect(v2)) as con:
            with con:
                con.executescript(ESQUEMA_BASE_SQL)
                con.executescript(ESQUEMA_V2_SQL)
                for nombre, definicion in (
                        ("barbero_id", "INTEGER REFERENCES barberos(id)"),
                        ("cliente", "TEXT"), ("metodo_pago", "TEXT"),
                        ("propina", "REAL NOT NULL DEFAULT 0")):
                    con.execute(f"ALTER TABLE ventas ADD COLUMN {nombre} {definicion}")
                con.execute("INSERT INTO barberos (nombre) VALUES ('Alexander')")
                con.execute("PRAGMA user_version = 2")

        db3 = ConexionDB(v2)
        db3.inicializar()
        bar3 = BarberoRepository(db3)
        alex3 = bar3.buscar_por_nombre("alexander")  # sin distinción de mayúsculas
        assert alex3 is not None and alex3.comision == 50.0, \
            "El barbero heredado de v2 debe quedar con comisión 50 %"
        assert len(MetodoPagoRepository(db3).listar()) == 4
        with closing(db3.abrir()) as con:
            cols = {f[1] for f in con.execute("PRAGMA table_info(barberos)")}
            assert "comision" in cols
            cols_v = {f[1] for f in con.execute("PRAGMA table_info(ventas)")}
            assert "cliente_id" in cols_v
            tablas = {f[0] for f in con.execute(
                "SELECT name FROM sqlite_master WHERE type='table'")}
            assert "adelantos" in tablas, "La migración debe crear adelantos"
            cols_c = {f[1] for f in con.execute("PRAGMA table_info(clientes)")}
            assert "fecha_nacimiento" in cols_c
            assert con.execute("PRAGMA user_version").fetchone()[0] == 5
        # La migración es idempotente: reinicializar no rompe nada
        db3.inicializar()
        assert bar3.obtener(alex3.id).comision == 50.0

        # 18. Limpieza de handles (Windows los bloquea)
        import gc; gc.collect()

    print("TODAS LAS PRUEBAS PASARON")
    print("  - Semillas: 7 ítems + dueño Juliano (50 %) + 4 métodos de pago;")
    print("    los empleados se agregan desde Administración")
    print("  - Registro con barbero, cliente, método canónico y propina")
    print("  - Validaciones rechazan sin guardar (método y cliente inválidos)")
    print("  - Admin: barberos con comisión 0-100, catálogo, movimientos")
    print("  - Adelantos: egreso de caja validado, eliminable; se reflejan en")
    print("    la hoja BALANCE y descuentan el NETO A COBRAR del barbero")
    print("  - Clientes: celular único, mail opcional, baja sin tocar historial")
    print("  - Clientes: NACIMIENTO opcional, normalizado DD/MM/AAAA y")
    print("    validado (formato real, nunca fecha futura)")
    print("  - Clientes: autocompletado por nombre/apellido/celular y")
    print("    detector de errores de tipeo (nombre + mail, celular distinto)")
    print("  - Moneda del sistema: símbolo en pantalla y en el Excel (G$)")
    print("  - Tema claro/oscuro persistido; modo claro pastel legible")
    print("  - Regresión Inventario: llamada por nombre + listas solo selección")
    print("  - Métodos de pago: agregar/eliminar, duplicados, último protegido")
    print("  - Exportación .xlsx: Registro + Dashboard (A LIQUIDAR) + hoja")
    print("    BALANCE (INGRESOS − EGRESOS) + hoja CLIENTES (con NACIMIENTO) +")
    print("    hoja por barbero con LIQUIDACIÓN, ADELANTOS y NETO A COBRAR")
    print("  - Migraciones automáticas v1 → v5 y v2 → v5 sin perder datos")


if __name__ == "__main__":
    main()
